from __future__ import annotations

from pathlib import Path
from shutil import copyfile

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image


ROOT = Path(__file__).resolve().parents[1]
REF = ROOT / "outputs" / "reference_style_reproduction"
GPT = ROOT / "outputs" / "gptimage2_diagrams"
SUPP = ROOT / "outputs" / "supplementary_figures"
OUT_CN = ROOT / "outputs" / "导师汇报版_图表设计与结果解读_NSCLC_MMFT.docx"
OUT_ASCII = ROOT / "outputs" / "advisor_figure_plan_nsclc_mmft.docx"


FONT = "Microsoft YaHei"
INK = RGBColor(31, 41, 55)
BLUE = RGBColor(31, 78, 121)
MUTED = RGBColor(85, 96, 110)
RED = RGBColor(153, 27, 27)
GREEN = RGBColor(53, 94, 59)


FIGURES = [
    {
        "label": "图1",
        "title": "模型总体性能",
        "path": REF / "Figure1_reference_style_performance.png",
        "one": "本图作为模型性能主图，用于展示公共测试集中不同模型的区分能力、预测分布和阈值相关指标。",
        "talk": "公共测试集结果显示，MMFT rank-refit 取得 AUC=0.868，较 Ridge logistic（AUC=0.854）略有提升；混淆矩阵和风险分布提示模型具备一定风险分层能力。",
        "caution": "后续需接入本院独立队列，补充外部验证 ROC、校准曲线和 DCA，以评价模型泛化能力。",
    },
    {
        "label": "图2",
        "title": "亚组稳定性",
        "path": REF / "Figure2_reference_style_subgroup_auc.png",
        "one": "本图用于评估模型在不同临床分层中的稳定性，减少模型仅在特定人群中有效的疑虑。",
        "talk": "按年龄、CEA、SUVmax、NLR、分期和性别分层后，多数亚组仍保持可接受的 AUC，提示模型预测信号并非完全依赖单一变量。",
        "caution": "部分亚组样本量较小，置信区间较宽；后续本院验证中应优先保留临床意义明确、样本量相对充足的亚组。",
    },
    {
        "label": "图3",
        "title": "个体解释和 SHAP",
        "path": REF / "Figure3_reference_style_case_interpretability.png",
        "one": "本图用于展示模型在个体病例层面的可解释性，说明预测结果主要由哪些变量推动。",
        "talk": "代表性病例的 SHAP 分解显示，融合风险 token、SUVmin、CEA、SUVmean 等变量贡献较高，与影像代谢负荷和血清标志物相关的研究假设一致。",
        "caution": "SHAP 结果用于解释模型预测依据，不能直接推断因果关系；后续可在本院验证队列中复核变量重要性排序。",
    },
    {
        "label": "图4",
        "title": "TCGA 生存曲线",
        "path": REF / "Figure4_reference_style_km_tcga_real.png",
        "one": "本图基于 TCGA-LUAD/LUSC 外部转录组与生存数据，为模型变量提供生物学层面的辅助解释。",
        "talk": "由 CEACAM5、SLC2A1、HK2、LDHA、TP63、KRT5 构建的 biological risk proxy 在 LUAD 中呈现较明显的生存分层，支持 CEA、糖酵解代谢和病理分化轴具有一定外部生物学依据。",
        "caution": "该分析属于干实验论证，不能替代临床外部验证；论文中应表述为机制支持，而非模型生存预测结果。",
    },
    {
        "label": "图5",
        "title": "临床路径图",
        "path": REF / "Figure5_reference_style_clinical_pathway.png",
        "one": "本图用于说明模型与胸外科术前评估流程之间的衔接关系。",
        "talk": "图中将术前影像、血液/炎症指标和 CEA 等常规信息整合为模型输入，输出患者层面的风险概率，并通过本院独立队列进行验证。",
        "caution": "该图定位为研究路径和临床转化设想，后续仍需以外部验证结果决定是否具备实际应用价值。",
    },
    {
        "label": "图6",
        "title": "队列与纳排流程",
        "path": REF / "Figure6_reference_style_participant_flow.png",
        "one": "本图明确公共建模队列、TCGA 干实验队列和后续本院验证队列之间的关系。",
        "talk": "公共 PET/CT-blood 队列用于模型开发和内部测试，TCGA-LUAD/LUSC 用于生物学论证，本院约 100 例病例计划作为锁定模型后的独立验证集。",
        "caution": "本院病例建议仅用于外部验证，不再参与模型调参，以保证验证结果的独立性。",
    },
    {
        "label": "图7",
        "title": "GPTimage2 研究总流程图",
        "path": GPT / "Figure7_GPTimage2_research_workflow_final.png",
        "one": "本图作为研究总流程图，用于概括数据来源、模型构建、结果评价和后续验证的完整技术路线。",
        "talk": "流程图分为数据、模型和分析三个层次，突出公共队列建模、TCGA 生物学论证和本院验证三个核心模块。",
        "caution": "正式成稿时可作为 graphical workflow 使用；图中文字和节点已进行后期校正，后续仅需根据最终研究结局微调输出端描述。",
    },
    {
        "label": "图8",
        "title": "GPTimage2 模型架构图",
        "path": GPT / "Figure9_GPTimage2_model_architecture_final.png",
        "one": "本图展示 MMFT 模型的输入分支、特征 token 化、Transformer 融合、风险输出和解释模块。",
        "talk": "模型输入包括临床信息、血液/CEA 指标、PET/CT 代谢特征和可选 CT 语义特征；Transformer encoder 用于学习模态间交互，risk token 汇总患者层面的融合风险表示。",
        "caution": "若本院验证队列无法获得 PET/CT 指标，可将 PET/CT 代谢分支设为可选模块，并同步构建仅含胸部 CT 与血液指标的简化版本。",
    },
    {
        "label": "图9",
        "title": "TCGA 生物信息学补强",
        "path": REF / "Figure8_TCGA_external_bioinfo_validation.png",
        "one": "本图从外部转录组角度补充模型变量的生物学合理性，是全文机制解释部分的重要图件。",
        "talk": "TCGA 分析显示，CEACAM5、糖酵解相关基因及鳞癌分化标志物在 LUAD 与 LUSC 间呈现方向性差异，并可与模型中的 CEA 和代谢特征形成呼应。",
        "caution": "该部分建议保持为辅助论证，不宜喧宾夺主；主线仍应落在术前多模态模型构建和本院外部验证。",
    },
]


TABLE_FIGURES = [
    {
        "label": "表1",
        "title": "公共队列基线特征",
        "path": REF / "Table1_baseline_characteristics_public_cohort.png",
        "one": "本表总结公共 PET/CT-blood 队列的基线特征和候选预测变量。",
        "talk": "队列中包含年龄、BMI、血常规、炎症指标、CEA 及 PET 代谢参数，与本课题拟纳入的术前变量体系基本一致。",
        "caution": "后续本院数据表建议尽量沿用相同变量命名和单位；同时需在方法部分明确 Class 0/1 的病理标签编码。",
    },
    {
        "label": "表2",
        "title": "模型性能表",
        "path": REF / "Table2_model_performance_public_test.png",
        "one": "本表列出公共测试集中不同模型的主要性能指标，便于结果部分直接引用。",
        "talk": "Ridge logistic 作为传统机器学习强基线，AUC 为 0.854；MMFT rank-refit 取得当前最佳表现，AUC 为 0.868，并在 accuracy、specificity 和 F1 上具有一定优势。",
        "caution": "正式论文中应同时报告 AUC、accuracy、sensitivity、specificity、PPV、NPV 和 F1，并结合校准曲线与 DCA 进行完整评价。",
    },
    {
        "label": "表3",
        "title": "消融与落地计划",
        "path": REF / "Table3_ablation_and_landing_plan.png",
        "one": "本表梳理不同模型配置和模块贡献，用于支撑模型设计的合理性。",
        "talk": "表中分别列出临床/血液、PET/CT 代谢特征、CEA、risk token 和 rank-refit 等模块，显示最终模型较传统基线存在小幅提升。",
        "caution": "本院验证结果完成前，相关条目应继续标注为 pending，避免将计划性分析写成已完成结果。",
    },
]


SUPPLEMENTARY = [
    {
        "label": "补充图总览",
        "title": "补充材料样式图",
        "path": SUPP / "Supplementary_Figure_Style_Sheet_NSCLC_MMFT.png",
        "one": "本图汇总补充材料的整体版式和内容安排，用于展示补充图体系的完整性。",
        "talk": "FigS1-FigS5 已基于现有公共数据、模型结果、SHAP 和 TCGA 数据生成样式图；FigS6 作为本院外部验证结果的预留模板。",
        "caution": "该总览图主要用于阶段性汇报，正式论文中可拆分为独立补充图，不建议作为单独论文图件提交。",
    },
    {
        "label": "FigS1",
        "title": "数据预处理与缺失值",
        "path": SUPP / "FigS1_data_preprocessing_missingness.png",
        "one": "本图展示数据预处理、缺失值概况、特征筛选流程及训练/测试划分。",
        "talk": "公共队列中关键临床、血液和 PET 代谢变量缺失率较低；特征筛选在交叉验证流程内完成，可降低测试集信息泄漏风险。",
        "caution": "本院队列纳入后，应重新绘制本院数据缺失率与变量保留流程，不能直接沿用公共队列的数据质量结论。",
    },
    {
        "label": "FigS2",
        "title": "模型消融实验",
        "path": SUPP / "FigS2_model_ablation.png",
        "one": "本图展示不同特征模块和建模策略的消融比较。",
        "talk": "clinical-blood、PET metabolic、radiomics、risk token 和 rank-refit 分别进行比较，有助于说明最终模型性能提升主要来自哪些模块。",
        "caution": "正式成稿时需区分 real benchmark、prior run 和 final run，避免将不同实验批次混写为同一轮严格消融。",
    },
    {
        "label": "FigS3",
        "title": "TCGA/GEO 关键基因方向热图",
        "path": SUPP / "FigS3_TCGA_GEO_key_gene_direction_heatmap.png",
        "one": "本图用于补充关键基因在外部表达队列中的方向性证据。",
        "talk": "当前版本基于真实 TCGA-LUAD/LUSC 数据生成，显示腺癌标志物、鳞癌标志物和糖酵解相关基因的表达方向差异。",
        "caution": "GEO 队列尚未正式接入前，论文图题建议写作 TCGA external validation；GEO 可作为后续扩展分析预留。",
    },
    {
        "label": "FigS4",
        "title": "校准曲线与 DCA",
        "path": SUPP / "FigS4_calibration_DCA.png",
        "one": "本图用于补充模型校准度和临床净获益评价。",
        "talk": "左侧校准曲线评估预测概率与观察比例的一致性，右侧 DCA 比较不同阈值下模型相对于 treat all 和 treat none 策略的净获益。",
        "caution": "公共测试集样本量有限，校准和 DCA 结果应作为初步评价；本院外部验证后需重新绘制并作为重点结果。",
    },
    {
        "label": "FigS5",
        "title": "SHAP 补充图",
        "path": SUPP / "FigS5_SHAP_supplement_composite.png",
        "one": "本图汇总 SHAP 解释性分析的四种常用展示形式。",
        "talk": "beeswarm 展示变量影响方向和离散程度，bar 展示平均绝对贡献，dependence 展示 CEA 与 SUVmax 相关的局部效应，waterfall 展示代表性病例的预测分解。",
        "caution": "SHAP 结果定位为模型解释性证据，后续需在本院验证队列中观察重要变量排序是否保持一致。",
    },
    {
        "label": "FigS6",
        "title": "本院验证队列模板",
        "path": SUPP / "FigS6_hospital_validation_template.png",
        "one": "本图为本院外部验证结果的预留版式。",
        "talk": "后续本院约 100 例病例收集完成后，可在该图中补入外部验证 ROC、校准曲线、DCA 和亚组稳定性分析。",
        "caution": "当前版本应继续明确标注 pending，待真实本院数据接入后再替换为正式结果图。",
    },
]


SHAP_SINGLE = [
    ("FigS5A", "SHAP beeswarm 单图", SUPP / "FigS5A_SHAP_beeswarm_highres.png"),
    ("FigS5B", "SHAP bar 单图", SUPP / "FigS5B_SHAP_bar_highres.png"),
    ("FigS5C", "SHAP dependence 单图", SUPP / "FigS5C_SHAP_dependence_CEA_highres.png"),
    ("FigS5D", "SHAP waterfall 单图", SUPP / "FigS5D_SHAP_waterfall_highres.png"),
]


def set_font(run, size: float | None = None, bold: bool = False, color: RGBColor | None = None) -> None:
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    if size is not None:
        run.font.size = Pt(size)
    run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_margins(section) -> None:
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    section.top_margin = Inches(0.42)
    section.bottom_margin = Inches(0.42)
    section.left_margin = Inches(0.48)
    section.right_margin = Inches(0.48)
    section.header_distance = Inches(0.25)
    section.footer_distance = Inches(0.25)


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    p = doc.add_paragraph()
    r = p.add_run(text)
    set_font(r, 16 if level == 1 else 12.5, True, BLUE if level == 1 else INK)
    p.paragraph_format.space_before = Pt(8 if level == 1 else 5)
    p.paragraph_format.space_after = Pt(4)


def add_para(doc: Document, text: str, size: float = 9.4, color: RGBColor = INK) -> None:
    p = doc.add_paragraph()
    r = p.add_run(text)
    set_font(r, size, False, color)
    p.paragraph_format.line_spacing = 1.16
    p.paragraph_format.space_after = Pt(3)


def add_labeled_para(doc: Document, label: str, text: str, color: RGBColor = BLUE) -> None:
    p = doc.add_paragraph()
    r1 = p.add_run(label)
    set_font(r1, 9.2, True, color)
    r2 = p.add_run(text)
    set_font(r2, 9.2, False, INK)
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(3)


def image_width(path: Path, max_width: float = 9.35, max_height: float = 4.7) -> float:
    try:
        with Image.open(path) as im:
            ratio = im.height / im.width
        return max(5.6, min(max_width, max_height / ratio))
    except Exception:
        return max_width


def add_image(doc: Document, path: Path, max_width: float = 9.35, max_height: float = 4.7) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if path.exists():
        p.add_run().add_picture(str(path), width=Inches(image_width(path, max_width, max_height)))
    else:
        r = p.add_run(f"图片缺失：{path}")
        set_font(r, 10, True, RED)
    p.paragraph_format.space_after = Pt(5)


def add_figure_page(doc: Document, item: dict, max_height: float = 4.55) -> None:
    doc.add_page_break()
    add_heading(doc, f"{item['label']}｜{item['title']}", level=1)
    add_image(doc, item["path"], max_width=9.45, max_height=max_height)
    add_labeled_para(doc, "图表目的：", item["one"], BLUE)
    add_labeled_para(doc, "主要解读：", item["talk"], GREEN)
    add_labeled_para(doc, "后续完善：", item["caution"], RED)


def add_cover(doc: Document) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("NSCLC-MMFT 课题图表设计与结果解读")
    set_font(r, 22, True, BLUE)
    p.paragraph_format.space_after = Pt(3)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("导师汇报版｜公共队列建模、TCGA 干实验与本院验证预案")
    set_font(r, 11, False, MUTED)
    p.paragraph_format.space_after = Pt(10)

    add_para(
        doc,
        "本文件用于课题进展汇报和方案讨论，整合现阶段已完成的公共队列建模结果、TCGA 生物信息学论证、GPTimage2 辅助绘制的流程/架构图，以及后续本院外部验证的补充图框架。",
        size=10,
    )
    add_para(
        doc,
        "整体研究主线建议保持为：公共 PET/CT-blood 队列用于模型构建，TCGA-LUAD/LUSC 用于生物学合理性论证，本院约 100 例病例作为后续独立验证队列。该设计能在保证可执行性的同时，补足模型解释性和外部验证逻辑。",
        size=10,
    )

    add_heading(doc, "建议汇报逻辑", level=1)
    rows = [
        ("研究问题", "围绕胸外科术前评估场景，探索血液/炎症指标与影像代谢特征联合建模的可行性。"),
        ("数据来源", "公共 PET/CT-blood 队列用于前期建模，TCGA 队列用于机制层面论证，本院病例用于后续独立验证。"),
        ("当前结果", "已完成模型性能、亚组稳定性、SHAP 解释、TCGA 生存/表达分析及补充图样式设计。"),
        ("待完善内容", "下一阶段重点是本院外部验证，并根据实际可获得变量锁定最终模型版本。"),
    ]
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    widths = [Inches(1.55), Inches(8.55)]
    for idx, h in enumerate(["环节", "怎么说"]):
        cell = table.rows[0].cells[idx]
        cell.text = h
        set_cell_shading(cell, "E8EEF5")
        for p in cell.paragraphs:
            for r in p.runs:
                set_font(r, 9, True, BLUE)
    for left, right in rows:
        cells = table.add_row().cells
        cells[0].text = left
        cells[1].text = right
        for i, cell in enumerate(cells):
            cell.width = widths[i]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cell.paragraphs:
                for r in p.runs:
                    set_font(r, 8.6, i == 0, INK)
    doc.add_paragraph()


def add_figure_index(doc: Document) -> None:
    doc.add_page_break()
    add_heading(doc, "图表清单", level=1)
    add_para(doc, "主图用于呈现研究主线和核心结果，补充图用于呈现数据处理、消融实验、校准、DCA 与解释性分析等方法学细节。正式成稿时可根据篇幅和期刊/学校要求，将部分图件调整至补充材料。", size=9.5)
    rows = []
    for item in FIGURES:
        rows.append((item["label"], item["title"], "主图"))
    for item in TABLE_FIGURES:
        rows.append((item["label"], item["title"], "表格图"))
    for item in SUPPLEMENTARY:
        rows.append((item["label"], item["title"], "补充图"))

    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for idx, h in enumerate(["编号", "标题", "用途"]):
        cell = table.rows[0].cells[idx]
        cell.text = h
        set_cell_shading(cell, "E8EEF5")
        for p in cell.paragraphs:
            for r in p.runs:
                set_font(r, 8.6, True, BLUE)
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            cells[idx].text = value
            for p in cells[idx].paragraphs:
                for r in p.runs:
                    set_font(r, 8.2, False, INK)


def add_shap_single_page(doc: Document) -> None:
    doc.add_page_break()
    add_heading(doc, "FigS5A-D｜SHAP 单图备用版", level=1)
    add_para(doc, "本页为 SHAP 解释性分析的单图备用版。后续可根据正文版面需要，选择合并展示或拆分为独立补充图。", size=9.2, color=MUTED)
    table = doc.add_table(rows=2, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for idx, (label, title, path) in enumerate(SHAP_SINGLE):
        cell = table.rows[idx // 2].cells[idx % 2]
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        cell.text = ""
        p = cell.paragraphs[0]
        r = p.add_run(f"{label}｜{title}")
        set_font(r, 8.3, True, BLUE)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if path.exists():
            p_img = cell.add_paragraph()
            p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_img.add_run().add_picture(str(path), width=Inches(4.6))
        else:
            p_missing = cell.add_paragraph()
            r_missing = p_missing.add_run(f"缺失：{path.name}")
            set_font(r_missing, 8, True, RED)


def build_doc() -> None:
    doc = Document()
    set_margins(doc.sections[0])

    styles = doc.styles
    styles["Normal"].font.name = FONT
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    styles["Normal"].font.size = Pt(9.5)

    add_cover(doc)
    add_figure_index(doc)

    for item in FIGURES:
        add_figure_page(doc, item)

    for item in TABLE_FIGURES:
        add_figure_page(doc, item, max_height=4.25)

    for item in SUPPLEMENTARY:
        add_figure_page(doc, item)
    add_shap_single_page(doc)

    doc.add_page_break()
    add_heading(doc, "最后提醒", level=1)
    add_para(doc, "当前图表体系已可用于阶段性汇报。下一步需要优先确认最终结局变量：若以病理亚型预测为主，现有公共库和 TCGA 论证链条较为完整；若转向淋巴结转移或复发预测，则本院数据必须提供相应金标准标签。", size=10)
    add_para(doc, "后面真正收本院数据时，建议先把字段模板锁定：年龄、性别、吸烟、血常规、NLR/PLR/SII、CEA、CT 语义特征、PET SUV 指标、最终病理标签。这样后面替换 FigS6 和外部验证表会快很多。", size=10)

    doc.save(OUT_CN)
    copyfile(OUT_CN, OUT_ASCII)
    print(OUT_CN)
    print(OUT_ASCII)


if __name__ == "__main__":
    build_doc()
