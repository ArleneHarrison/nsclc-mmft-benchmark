from __future__ import annotations

import math
import textwrap
from pathlib import Path

import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "outputs"
FIG_DIR = OUT_DIR / "figures_style"
RESULT_DIR = ROOT / "server_results" / "petct_blood_benchmark"
DOCX_OUT = OUT_DIR / "NSCLC_MMFT_图表思路_中文.docx"


def roc_curve_np(y_true, y_score):
    y_true = np.asarray(y_true).astype(int)
    y_score = np.asarray(y_score).astype(float)
    order = np.argsort(-y_score)
    y_true = y_true[order]
    y_score = y_score[order]
    positives = max(1, int((y_true == 1).sum()))
    negatives = max(1, int((y_true == 0).sum()))
    tps = np.cumsum(y_true == 1)
    fps = np.cumsum(y_true == 0)
    distinct = np.r_[True, y_score[1:] != y_score[:-1]]
    tpr = np.r_[0, tps[distinct] / positives, 1]
    fpr = np.r_[0, fps[distinct] / negatives, 1]
    return fpr, tpr


def setup_font() -> None:
    plt.rcParams["font.sans-serif"] = ["Microsoft YaHei", "SimHei", "Arial Unicode MS", "DejaVu Sans"]
    plt.rcParams["axes.unicode_minus"] = False
    plt.rcParams["figure.facecolor"] = "white"
    plt.rcParams["savefig.facecolor"] = "white"


def add_watermark(ax, text="STYLE ONLY / 待真实数据替换") -> None:
    ax.text(
        0.5,
        0.5,
        text,
        transform=ax.transAxes,
        ha="center",
        va="center",
        fontsize=18,
        color="#B00020",
        alpha=0.16,
        rotation=18,
        weight="bold",
    )


def draw_workflow() -> Path:
    path = FIG_DIR / "Figure1_workflow.png"
    fig, ax = plt.subplots(figsize=(13, 7), dpi=180)
    ax.axis("off")

    boxes = [
        ("Public cohort\n255 NSCLC", 0.05, 0.68, "#E8F1FA"),
        ("Feature blocks\nblood | PET | radiomics", 0.29, 0.68, "#E8F1FA"),
        ("Benchmark\nridge/LASSO/tree", 0.53, 0.68, "#E8F1FA"),
        ("MMFT rank-refit\nrisk token + rank loss", 0.77, 0.68, "#DFF3EA"),
        ("Prediction\nLUAD vs LUSC", 0.77, 0.40, "#FEF3C7"),
        ("Dry biology\nTCGA/GEO proof", 0.53, 0.40, "#FCE7F3"),
        ("Hospital validation\nblood + CT/PET", 0.29, 0.40, "#F3F4F6"),
        ("Final story\nmetabolism + CEA + histology", 0.41, 0.14, "#E0F2FE"),
    ]

    def box(label, x, y, color):
        rect = plt.Rectangle((x, y), 0.18, 0.15, fc=color, ec="#334155", lw=1.5, joinstyle="round")
        ax.add_patch(rect)
        ax.text(x + 0.09, y + 0.075, label, ha="center", va="center", fontsize=9.5, color="#111827")

    for b in boxes:
        box(*b)

    arrows = [
        ((0.23, 0.755), (0.29, 0.755)),
        ((0.47, 0.755), (0.53, 0.755)),
        ((0.71, 0.755), (0.77, 0.755)),
        ((0.86, 0.68), (0.86, 0.55)),
        ((0.77, 0.475), (0.71, 0.475)),
        ((0.53, 0.475), (0.47, 0.475)),
        ((0.38, 0.40), (0.46, 0.29)),
        ((0.62, 0.40), (0.55, 0.29)),
        ((0.86, 0.40), (0.58, 0.25)),
    ]
    for start, end in arrows:
        ax.annotate("", xy=end, xytext=start, arrowprops=dict(arrowstyle="->", lw=1.6, color="#334155"))

    ax.text(0.5, 0.95, "Study workflow: modeling mainline plus biological and clinical validation branches", ha="center", fontsize=15, weight="bold", color="#0F172A")
    ax.text(0.5, 0.035, "Dry bioinformatics panels are style templates and must be replaced by real TCGA/GEO outputs before submission.", ha="center", fontsize=8.8, color="#64748B")
    fig.tight_layout()
    fig.savefig(path, bbox_inches="tight")
    plt.close(fig)
    return path


def draw_performance() -> Path:
    path = FIG_DIR / "Figure2_model_performance.png"
    ridge = pd.read_csv(RESULT_DIR / "best_model_test_predictions.csv")
    mmft = pd.read_csv(RESULT_DIR / "mmft_rank_refit_predictions.csv")
    perf = pd.read_csv(RESULT_DIR / "mmft_rank_refit_results.csv").iloc[0].to_dict()

    fig, axes = plt.subplots(1, 3, figsize=(14, 4.4), dpi=180)
    for df, label, color, ax_auc in [
        (ridge, "Ridge logistic (AUC=0.854)", "#2563EB", None),
        (mmft, "MMFT rank-refit (AUC=0.868)", "#DC2626", None),
    ]:
        fpr, tpr = roc_curve_np(df["histology"], df["predicted_probability"])
        axes[0].plot(fpr, tpr, lw=2.2, color=color, label=label)
    axes[0].plot([0, 1], [0, 1], "--", color="#94A3B8", lw=1)
    axes[0].set_title("A. ROC comparison", weight="bold")
    axes[0].set_xlabel("False positive rate")
    axes[0].set_ylabel("True positive rate")
    axes[0].legend(frameon=False, fontsize=8)
    axes[0].grid(alpha=0.25)

    metrics = ["auc", "accuracy", "sensitivity", "specificity"]
    mmft_vals = [perf[m] for m in metrics]
    ridge_vals = [0.8539944903581267, 0.7402597402597403, 0.8636363636363636, 0.5757575757575758]
    x = np.arange(len(metrics))
    axes[1].bar(x - 0.18, ridge_vals, width=0.36, color="#93C5FD", label="Ridge")
    axes[1].bar(x + 0.18, mmft_vals, width=0.36, color="#FCA5A5", label="MMFT")
    axes[1].set_ylim(0, 1)
    axes[1].set_xticks(x)
    axes[1].set_xticklabels(["AUC", "ACC", "SEN", "SPE"], fontsize=8)
    axes[1].set_title("B. Test-set metrics", weight="bold")
    axes[1].legend(frameon=False, fontsize=8)
    axes[1].grid(axis="y", alpha=0.25)

    ranked = mmft.sort_values("predicted_probability").reset_index(drop=True)
    colors = np.where(ranked["histology"].to_numpy() == 1, "#DC2626", "#2563EB")
    axes[2].scatter(np.arange(len(ranked)), ranked["predicted_probability"], c=colors, s=18, alpha=0.82)
    axes[2].set_title("C. Individual risk distribution", weight="bold")
    axes[2].set_xlabel("Patients ordered by predicted probability")
    axes[2].set_ylabel("MMFT probability")
    axes[2].grid(alpha=0.25)
    axes[2].text(0.02, 0.94, "blue=class 0, red=class 1", transform=axes[2].transAxes, fontsize=8, color="#475569")

    fig.suptitle("Model performance based on existing public-data experiment", fontsize=13, weight="bold", y=1.03)
    fig.tight_layout()
    fig.savefig(path, bbox_inches="tight")
    plt.close(fig)
    return path


def draw_architecture() -> Path:
    path = FIG_DIR / "Figure3_mmft_architecture.png"
    fig, ax = plt.subplots(figsize=(13, 6.6), dpi=180)
    ax.axis("off")
    ax.text(0.5, 0.95, "MMFT-Transformer architecture", ha="center", fontsize=16, weight="bold", color="#0F172A")

    groups = [
        ("Clinical-blood tokens\nsex, CEA, blood routine", 0.05, 0.70, "#E0F2FE"),
        ("PET metabolic tokens\nSUVmean, SUVmax, SUVmin", 0.05, 0.48, "#DCFCE7"),
        ("Radiomics tokens\noptional filtered CT/PET features", 0.05, 0.26, "#F3E8FF"),
        ("Ridge risk token\nstrong linear prior", 0.05, 0.08, "#FEF3C7"),
    ]
    for label, x, y, color in groups:
        ax.add_patch(plt.Rectangle((x, y), 0.22, 0.13, fc=color, ec="#334155", lw=1.2))
        ax.text(x + 0.11, y + 0.065, label, ha="center", va="center", fontsize=9)

    layers = [
        ("Value projection\n+ feature ID\n+ modality ID", 0.36, 0.50, "#F8FAFC"),
        ("Transformer encoder\nself-attention across tokens", 0.55, 0.50, "#E8F1FA"),
        ("Gated modality pooling\n+ attention pooling", 0.74, 0.50, "#DFF3EA"),
        ("Rank-refit head\nBCE + pairwise AUC loss", 0.74, 0.22, "#FEE2E2"),
        ("Histology probability\nLUAD vs LUSC", 0.55, 0.22, "#FEF3C7"),
    ]
    for label, x, y, color in layers:
        ax.add_patch(plt.Rectangle((x, y), 0.16, 0.16, fc=color, ec="#334155", lw=1.3))
        ax.text(x + 0.08, y + 0.08, label, ha="center", va="center", fontsize=9)

    for y in [0.765, 0.545, 0.325, 0.145]:
        ax.annotate("", xy=(0.36, 0.58), xytext=(0.27, y), arrowprops=dict(arrowstyle="->", color="#475569", lw=1.4))
    for start, end in [((0.52, 0.58), (0.55, 0.58)), ((0.71, 0.58), (0.74, 0.58)), ((0.82, 0.50), (0.82, 0.38)), ((0.74, 0.30), (0.71, 0.30))]:
        ax.annotate("", xy=end, xytext=start, arrowprops=dict(arrowstyle="->", color="#475569", lw=1.4))
    ax.text(0.5, 0.02, "Design intent: use ridge risk as a stable prior, then let Transformer learn only the residual nonlinear ordering signal.", ha="center", fontsize=9, color="#64748B")
    fig.tight_layout()
    fig.savefig(path, bbox_inches="tight")
    plt.close(fig)
    return path


def draw_bioinformatics_panel() -> Path:
    path = FIG_DIR / "Figure4_bioinformatics_template.png"
    rng = np.random.default_rng(20260701)
    fig, axes = plt.subplots(2, 2, figsize=(13, 9), dpi=180)

    n = 900
    logfc = rng.normal(0, 1.2, n)
    p = np.clip(rng.beta(0.7, 8, n), 1e-8, 1)
    neglogp = -np.log10(p)
    sig = (np.abs(logfc) > 1) & (neglogp > 2)
    colors = np.where((logfc > 1) & (neglogp > 2), "#DC2626", np.where((logfc < -1) & (neglogp > 2), "#2563EB", "#CBD5E1"))
    ax = axes[0, 0]
    ax.scatter(logfc, neglogp, c=colors, s=10, alpha=0.72, edgecolor="none")
    ax.axvline(1, ls="--", color="#94A3B8", lw=1)
    ax.axvline(-1, ls="--", color="#94A3B8", lw=1)
    ax.axhline(2, ls="--", color="#94A3B8", lw=1)
    ax.set_title("A. TCGA LUAD vs LUSC volcano template", weight="bold")
    ax.set_xlabel("log2FC")
    ax.set_ylabel("-log10(P)")
    for gene, x, y in [("CEACAM5", 1.8, 5.2), ("NKX2-1", 1.3, 4.2), ("TP63", -1.7, 5.0), ("KRT5", -2.0, 4.7)]:
        ax.text(x, y, gene, fontsize=8, weight="bold")
    add_watermark(ax)

    ax = axes[0, 1]
    pathways = ["CEA/CEACAM", "Glycolysis", "Hypoxia", "Keratinization", "EMT", "Immune checkpoint", "OXPHOS"]
    nes = np.array([1.7, 1.45, 1.25, -1.9, -0.8, 0.7, -1.2])
    q = np.array([0.002, 0.006, 0.018, 0.001, 0.05, 0.08, 0.03])
    y = np.arange(len(pathways))
    ax.scatter(nes, y, s=(1 - q) * 240, c=nes, cmap="coolwarm", vmin=-2, vmax=2, edgecolor="#334155")
    ax.axvline(0, color="#94A3B8", lw=1)
    ax.set_yticks(y)
    ax.set_yticklabels(pathways, fontsize=8)
    ax.set_xlabel("NES")
    ax.set_title("B. Pathway enrichment template", weight="bold")
    add_watermark(ax)

    ax = axes[1, 0]
    genes = ["CEACAM5", "SLC2A1", "HK2", "LDHA", "TP63", "KRT5", "NKX2-1"]
    luad = rng.normal([7.2, 5.1, 5.5, 6.0, 3.0, 2.5, 7.5], 0.55, (60, len(genes)))
    lusc = rng.normal([4.2, 6.2, 6.4, 7.0, 7.3, 7.5, 2.5], 0.55, (60, len(genes)))
    positions = np.arange(len(genes))
    ax.boxplot(luad, positions=positions - 0.18, widths=0.28, patch_artist=True, boxprops=dict(facecolor="#93C5FD", color="#2563EB"), medianprops=dict(color="#1E3A8A"))
    ax.boxplot(lusc, positions=positions + 0.18, widths=0.28, patch_artist=True, boxprops=dict(facecolor="#FCA5A5", color="#DC2626"), medianprops=dict(color="#7F1D1D"))
    ax.set_xticks(positions)
    ax.set_xticklabels(genes, rotation=35, ha="right", fontsize=8)
    ax.set_ylabel("Expression")
    ax.set_title("C. Marker expression template", weight="bold")
    ax.text(0.02, 0.94, "blue=LUAD, red=LUSC", transform=ax.transAxes, fontsize=8)
    add_watermark(ax)

    ax = axes[1, 1]
    rows = ["Glycolysis", "Hypoxia", "CEACAM", "Neutrophil", "Macrophage", "T cell", "Stroma", "EMT"]
    cols = ["LUAD-low SUV", "LUAD-high SUV", "LUSC-low SUV", "LUSC-high SUV"]
    mat = np.array([
        [-0.5, 0.3, 0.2, 1.2],
        [-0.4, 0.4, 0.3, 1.1],
        [1.1, 1.3, -0.4, -0.2],
        [-0.3, 0.1, 0.5, 1.1],
        [-0.2, 0.1, 0.4, 0.9],
        [0.2, 0.1, -0.1, -0.5],
        [-0.3, 0.0, 0.3, 0.8],
        [-0.2, 0.2, 0.3, 1.0],
    ])
    im = ax.imshow(mat, cmap="RdBu_r", vmin=-1.5, vmax=1.5, aspect="auto")
    ax.set_xticks(np.arange(len(cols)))
    ax.set_xticklabels(cols, rotation=30, ha="right", fontsize=8)
    ax.set_yticks(np.arange(len(rows)))
    ax.set_yticklabels(rows, fontsize=8)
    ax.set_title("D. Metabolism-immune heatmap template", weight="bold")
    fig.colorbar(im, ax=ax, fraction=0.046, pad=0.02, label="z-score")
    add_watermark(ax)

    fig.suptitle("Bioinformatics dry-lab proof template: transcriptome, pathway and microenvironment rationale", fontsize=14, weight="bold")
    fig.tight_layout(rect=[0, 0, 1, 0.96])
    fig.savefig(path, bbox_inches="tight")
    plt.close(fig)
    return path


def draw_validation_roadmap() -> Path:
    path = FIG_DIR / "Figure5_validation_roadmap.png"
    fig, ax = plt.subplots(figsize=(13, 6), dpi=180)
    ax.axis("off")
    stages = [
        ("1. Public model\nPET/CT + blood\nn=255", "#E0F2FE"),
        ("2. Dry biology\nTCGA/GEO\npathway + immune", "#FCE7F3"),
        ("3. Hospital validation\n~100 cases\nroutine blood + CT/PET", "#DCFCE7"),
        ("4. Robustness checks\ncalibration + DCA\nablation + subgroup", "#FEF3C7"),
        ("5. Thesis output\ninterpretable model\nbiological rationale", "#F3F4F6"),
    ]
    xs = np.linspace(0.06, 0.82, len(stages))
    for i, ((label, color), x) in enumerate(zip(stages, xs)):
        ax.add_patch(plt.Rectangle((x, 0.45), 0.15, 0.22, fc=color, ec="#334155", lw=1.4))
        ax.text(x + 0.075, 0.56, label, ha="center", va="center", fontsize=9)
        if i < len(stages) - 1:
            ax.annotate("", xy=(xs[i + 1], 0.56), xytext=(x + 0.15, 0.56), arrowprops=dict(arrowstyle="->", lw=1.5, color="#334155"))

    ax.text(0.5, 0.88, "External validation and thesis landing roadmap", ha="center", fontsize=15, weight="bold")
    ax.text(0.06, 0.25, "Minimum hospital variables:", fontsize=10, weight="bold", color="#0F172A")
    ax.text(0.06, 0.18, "age, sex, pathology, CEA, WBC/NEU/LYM/PLT, NLR/dNLR, CT semantic signs, PET SUV if available, final histology", fontsize=9, color="#334155")
    ax.text(0.06, 0.08, "Critical check: report public-data result and hospital external validation separately; avoid presenting exploratory test-tuned AUC as unbiased performance.", fontsize=9, color="#9B1C1C")
    fig.tight_layout()
    fig.savefig(path, bbox_inches="tight")
    plt.close(fig)
    return path


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text: str, bold: bool = False) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if len(text) < 18 else WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(9)
    run.bold = bold


def style_table(table, header_fill="E8EEF5") -> None:
    table.autofit = False
    for row_i, row in enumerate(table.rows):
        for cell in row.cells:
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.1
            if row_i == 0:
                set_cell_shading(cell, header_fill)
                for p in cell.paragraphs:
                    for r in p.runs:
                        r.bold = True


def set_doc_styles(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing = 1.18
    normal.paragraph_format.space_after = Pt(6)

    for name, size, color, before, after in [
        ("Heading 1", 16, "1F4D78", 14, 8),
        ("Heading 2", 13, "2E74B5", 12, 6),
        ("Heading 3", 11.5, "1F4D78", 8, 4),
    ]:
        st = styles[name]
        st.font.name = "Microsoft YaHei"
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        st.font.size = Pt(size)
        st.font.color.rgb = RGBColor.from_string(color)
        st.font.bold = True
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)

    if "Compact" not in styles:
        styles.add_style("Compact", 1)
    compact = styles["Compact"]
    compact.font.name = "Microsoft YaHei"
    compact._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    compact.font.size = Pt(9.5)
    compact.paragraph_format.left_indent = Cm(0.45)
    compact.paragraph_format.space_after = Pt(3)
    compact.paragraph_format.line_spacing = 1.12

    if "Block Text" in styles:
        block = styles["Block Text"]
        block.font.name = "Microsoft YaHei"
        block._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        block.font.size = Pt(10)
        block.paragraph_format.left_indent = Cm(0.35)
        block.paragraph_format.right_indent = Cm(0.15)
        block.paragraph_format.space_before = Pt(4)
        block.paragraph_format.space_after = Pt(6)


def add_para(doc, text, style=None, bold_label=None):
    p = doc.add_paragraph(style=style)
    if bold_label and text.startswith(bold_label):
        run = p.add_run(bold_label)
        run.bold = True
        rest = text[len(bold_label):]
        p.add_run(rest)
    else:
        p.add_run(text)
    return p


def add_figure_section(doc, heading: str, image_path: Path, bullets: list[str]) -> None:
    doc.add_heading(heading, level=3)
    doc.add_picture(str(image_path), width=Inches(6.5))
    last = doc.paragraphs[-1]
    last.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for item in bullets:
        add_para(doc, item, style="Compact", bold_label=item.split(":")[0] + ":" if ":" in item else None)


def build_docx(figs: dict[str, Path]) -> None:
    doc = Document()
    set_doc_styles(doc)

    doc.add_heading("全文图表思路解读（中文）", level=1)
    doc.add_heading("文章标题", level=2)
    add_para(doc, "英文标题: A Multimodal Gated Feature Transformer Integrating Blood Markers and PET/CT Metabolic Features for Predicting Histological Subtype in Non-Small Cell Lung Cancer", style="Body Text")
    add_para(doc, "中文标题（建议）: 基于血液炎症指标及 PET/CT 代谢特征预测非小细胞肺癌病理亚型的多模态门控 Transformer 模型构建与生物信息学论证", style="Body Text")
    add_para(doc, "简明题目: Multimodal PET/CT-blood Transformer in NSCLC", style="Body Text")

    doc.add_heading("文章主要结果与发现（摘要）", level=2)
    add_para(doc, "背景与目的: 非小细胞肺癌腺癌与鳞癌在影像代谢、血清肿瘤标志物及分子通路层面具有不同表型。本研究以公开 PET/CT-血液指标数据为建模基础，构建可解释的多模态模型，并用干实验补充其生物学合理性。", style="Body Text")
    add_para(doc, "主要结果与发现:", style="Body Text")
    summary_items = [
        "公共库建模: 纳入 PLOS One 2024 公开队列 255 例 NSCLC，变量包括临床信息、血液指标、CEA、PET 代谢参数以及 CT/PET radiomics，结局为病理亚型。",
        "传统模型基线: clinical-blood-metabolic 特征组合表现最好，ridge logistic 测试集 AUC 0.854，优于单纯 radiomics 特征。",
        "Transformer 迭代: 普通 FT-Transformer AUC 0.815，普通 MMFT 约 0.821；诊断发现小样本、高维 radiomics 噪声和验证集不稳定是主要瓶颈。",
        "最终探索模型: 采用 ridge risk token + modality embedding + gated attention + pairwise ranking loss 的 MMFT rank-refit，测试集 AUC 0.868，准确率 0.792，敏感度 0.818，特异度 0.758。",
        "模型可解释性: SHAP 分析显示 ridge risk token 贡献最大，其次为 SUVmin、CEA、SUVmean、gender 和 SUVmax，提示模型主要由线性综合风险、葡萄糖代谢强度和 CEA 共同驱动。",
        "生物学解释: 最终变量集中在 gender、CEA、SUVmean/SUVmax/SUVmin 与 ridge risk token，提示模型主要捕捉“肿瘤标志物-葡萄糖摄取代谢-组织学表型”轴。",
        "干实验补强: 建议用 TCGA-LUAD/LUSC 与 GEO 队列验证 CEACAM5/CEA 轴、糖酵解-缺氧通路、鳞癌角化程序、免疫微环境差异，并把 PET 代谢特征与转录组糖酵解 proxy score 建立平行证据链。",
        "转化路径: 后续本院约 100 例作为外部验证队列，优先收集 CEA、WBC/NEU/LYM/PLT、NLR/dNLR、CT 语义特征、PET SUV（如有）和最终病理亚型。",
    ]
    for item in summary_items:
        add_para(doc, item, style="Compact")
    add_para(doc, "结论: 在当前公共数据中，传统正则化模型仍是最稳基线；经过面向小样本设计的 MMFT rank-refit 可获得更高探索性 AUC。论文写作时应将 0.868 标注为探索性终训结果，并以后续本院外部验证作为最终可信度来源。", style="Body Text")

    doc.add_heading("数据集总览（当前已用 + 建议补充）", level=2)
    table = doc.add_table(rows=1, cols=5)
    hdr = table.rows[0].cells
    for cell, text in zip(hdr, ["数据源", "类型", "样本/变量", "当前状态", "在文中的角色"]):
        set_cell_text(cell, text, bold=True)
    rows = [
        ["PLOS One 2024 PET/CT radiomics", "公开个体级数据", "255 例 NSCLC；血液、CEA、PET 代谢、CT/PET radiomics", "已下载并建模", "主建模队列"],
        ["NSCLC-Radiogenomics", "TCIA 影像-临床", "CT 语义特征；病理 N 分期有限", "已初步建模", "可作为备选胸外科 LN 方向"],
        ["TCGA-LUAD/LUSC", "转录组/临床", "RNA-seq；组织学、通路、免疫浸润", "待下载真实分析", "干实验机制论证"],
        ["GEO LUAD/LUSC 队列", "外部表达芯片/RNA-seq", "CEACAM5、糖酵解、鳞癌标志物", "待筛选", "方向复现"],
        ["本院队列", "回顾性外部验证", "约 100 例；血常规/CEA/CT/PET/病理", "待收集", "最终外部验证"],
    ]
    for row in rows:
        cells = table.add_row().cells
        for cell, text in zip(cells, row):
            set_cell_text(cell, text)
    style_table(table)

    doc.add_heading("一、整体思路（一句话主线）", level=2)
    add_para(doc, "非小细胞肺癌病理亚型不是单一影像问题，而是由肿瘤标志物表达、葡萄糖摄取代谢、组织学分化和宿主炎症共同塑造的多模态表型；因此本文用血液/CEA + PET 代谢 + 可选 radiomics 构建预测模型，并用 TCGA/GEO 生物信息学从 CEACAM5-糖酵解-免疫微环境三条线解释模型为什么可行。", style="Body Text")
    add_para(doc, "逻辑链: 公共 PET/CT-血液数据 → 传统模型基线 → MMFT-Transformer 结构优化 → 性能比较 → 模型变量解释 → TCGA/GEO 干实验论证 → 本院外部验证 → 临床转化边界。", style="Body Text")

    doc.add_heading("二、逐图解读", level=2)
    add_figure_section(
        doc,
        "Figure 1 | 研究流程图",
        figs["workflow"],
        [
            "展示: 从公开 PET/CT-血液队列出发，依次完成特征分块、传统机器学习基线、MMFT-Transformer 优化、生物信息学干实验、本院外部验证。",
            "关键数字: 主数据集 255 例；传统 ridge AUC 0.854；MMFT rank-refit 探索性 AUC 0.868。",
            "作用: 让读者一眼看到全文不是单纯调模型，而是“模型-解释-验证”闭环。",
            "说明: 图中干实验部分为样式图，真实 TCGA/GEO 结果后期替换。",
        ],
    )
    add_figure_section(
        doc,
        "Figure 2 | 模型表现与个体风险分布",
        figs["performance"],
        [
            "展示: Ridge 与 MMFT 的 ROC 对比、AUC/准确率/敏感度/特异度柱状图、个体预测概率分布。",
            "真实数字: ridge logistic AUC 0.854；MMFT rank-refit AUC 0.868，准确率 0.792，敏感度 0.818，特异度 0.758。",
            "作用: 证明最终模型相对强基线有一定提升，同时保留传统模型作为稳健参照。",
            "说明: 0.868 属探索性终训测试结果，投稿/毕业论文中建议以外部验证作为最终主结论。",
        ],
    )
    add_figure_section(
        doc,
        "Figure 3 | MMFT-Transformer 模型结构图",
        figs["architecture"],
        [
            "展示: 临床-血液 token、PET 代谢 token、可选 radiomics token 与 ridge risk token 进入同一 Transformer；经 modality embedding、gated pooling 和 rank-refit head 输出病理亚型概率。",
            "关键设计: ridge risk token 提供稳定线性先验，Transformer 只学习非线性残差信号；pairwise ranking loss 直接优化 AUC 排序。",
            "作用: 回应“为什么不是普通深度学习堆模型”，突出本模型专门为小样本表格/影像组学数据设计。",
            "说明: 可在方法中命名为 MMFT rank-refit model。",
        ],
    )
    add_figure_section(
        doc,
        "Figure 4 | 生物信息学干实验论证图（样式图）",
        figs["bioinfo"],
        [
            "展示: TCGA LUAD vs LUSC 火山图、通路富集气泡图、关键基因表达箱线图、代谢-免疫热图。",
            "待补真实分析: 差异基因、Hallmark/KEGG GSEA、CEACAM5/SLC2A1/HK2/LDHA/TP63/KRT5/NKX2-1 表达、xCell/CIBERSORT 免疫浸润。",
            "作用: 为模型变量提供生物学解释——CEA/CEACAM5 指向腺癌分泌表型，SUV 指向糖酵解代谢，鳞癌可由 TP63/KRT5/角化程序作为对照轴。",
            "说明: 当前图为模拟样式图，已加水印，后期必须替换为真实 TCGA/GEO 结果。",
        ],
    )
    add_figure_section(
        doc,
        "Figure 5 | 本院外部验证与论文落地路线图",
        figs["roadmap"],
        [
            "展示: 公共建模、干实验论证、本院外部验证、校准/DCA/亚组、论文输出五步路线。",
            "关键采集项: CEA、血常规、NLR/dNLR、CT 语义特征、PET SUV（如有）、最终病理亚型。",
            "作用: 把当前公共库模型转化为可毕业、可验证的临床预测模型研究。",
            "说明: 如果本院没有 PET，可退化为 CEA + 血常规 + 术前 CT 语义特征的外部验证。",
        ],
    )
    if "shap" in figs:
        add_figure_section(
            doc,
            "Figure 6 | SHAP 模型可解释性分析",
            figs["shap"],
            [
                "展示: SHAP 蜂群图、平均绝对 SHAP 重要性、CEA 依赖图和代表性高风险病例 waterfall 图。",
                "真实数字: mean(|SHAP|) 排名为 ridge_risk_token 0.157、SUVmin 0.049、CEA 0.035、SUVmean 0.026、gender 0.026、SUVmax 0.014。",
                "作用: 解释 MMFT rank-refit 的预测主要由强线性风险先验驱动，同时由 PET 代谢参数和 CEA 提供增量信息。",
                "说明: ridge_risk_token 是由 gender、CEA、SUVmean、SUVmax、SUVmin 线性组合得到的风险先验，因此解释时应同时报告原始变量 SHAP 和风险 token 的模型层作用。",
            ],
        )

    doc.add_heading("三、建议新增干实验（生物信息学）", level=2)
    dry_table = doc.add_table(rows=1, cols=4)
    for cell, text in zip(dry_table.rows[0].cells, ["干实验", "具体做法", "预期图形", "与模型的对应关系"]):
        set_cell_text(cell, text, bold=True)
    dry_rows = [
        ["TCGA 组织学差异分析", "LUAD vs LUSC 差异基因；标注 CEACAM5、TP63、KRT5、NKX2-1", "火山图 + 标志基因箱线图", "解释 CEA 与组织学表型"],
        ["糖酵解/PET proxy", "GSVA 计算 glycolysis、hypoxia、FDG-uptake proxy score", "GSEA 曲线/气泡图", "解释 SUVmean/SUVmax/SUVmin"],
        ["免疫炎症微环境", "xCell/CIBERSORT/MCP-counter 估计中性粒、巨噬、T 细胞等", "免疫热图/相关矩阵", "连接血常规/NLR 与肿瘤环境"],
        ["模型变量-通路桥接", "按 CEACAM5 高低或 histology risk score 分组比较通路", "通路热图/森林图", "把模型输出转成机制解释"],
        ["外部表达队列复现", "GEO 中验证 CEACAM5、糖酵解和鳞癌标志物方向", "跨队列 AUC/方向热图", "增强泛化可信度"],
    ]
    for row in dry_rows:
        cells = dry_table.add_row().cells
        for cell, text in zip(cells, row):
            set_cell_text(cell, text)
    style_table(dry_table)

    doc.add_heading("四、补充图建议", level=2)
    supps = [
        "FigS1 | 数据预处理与缺失值图: 展示各变量缺失率、特征筛选流程、训练/测试划分。",
        "FigS2 | 模型消融实验: clinical-blood、PET metabolic、radiomics、risk token、rank loss 分别比较。",
        "FigS3 | TCGA/GEO 真实干实验补充: 外部队列关键基因方向热图。",
        "FigS4 | 校准与 DCA: 对 ridge 和 MMFT 分别做 calibration curve、decision curve。",
        "FigS5 | SHAP 补充图: 分别展示 beeswarm、bar、dependence、waterfall 的高清单图。",
        "FigS6 | 本院验证队列分析: 外部验证 ROC、校准、DCA 和亚组稳定性。",
    ]
    for item in supps:
        add_para(doc, item, style="Compact")

    doc.add_heading("五、自检清单", level=2)
    checks = [
        "主题切题: 模型围绕 NSCLC 病理亚型、术前 PET/CT/血液指标展开，符合胸外科/肺癌方向。",
        "新颖性: 不是单纯 radiomics，而是 blood-PET/CT-MMFT + ridge risk token + 生物信息学解释。",
        "结果边界: 已明确 0.868 为探索性终训结果，不能替代外部验证。",
        "干实验边界: 生信图当前为样式图，已标注待真实 TCGA/GEO 数据替换。",
        "可落地性: 本院 100 例可优先验证 CEA、血常规、NLR/dNLR、CT 语义特征和病理亚型。",
        "论文风险: 如果本院无 PET，应把题目收窄为 CT 语义 + 血液指标；如果有 PET，则保留 PET/CT 代谢特征作为亮点。",
    ]
    for item in checks:
        add_para(doc, "□ " + item, style="Compact")

    doc.add_heading("六、写作建议", level=2)
    add_para(doc, "建议主文写法: 以 ridge logistic 作为稳健基线，以 MMFT rank-refit 作为探索性深度学习增强模型；生物信息学部分不直接宣称因果机制，而写为“外部转录组层面的生物学合理性论证”。", style="Body Text")
    add_para(doc, "最稳妥论文定位: “术前血液指标及 PET/CT 代谢特征预测 NSCLC 病理亚型的机器学习模型构建及生物信息学论证”。", style="Body Text")

    doc.save(DOCX_OUT)


def main() -> None:
    OUT_DIR.mkdir(exist_ok=True)
    FIG_DIR.mkdir(parents=True, exist_ok=True)
    setup_font()
    figs = {
        "workflow": draw_workflow(),
        "performance": draw_performance(),
        "architecture": draw_architecture(),
        "bioinfo": draw_bioinformatics_panel(),
        "roadmap": draw_validation_roadmap(),
    }
    shap_path = FIG_DIR / "Figure6_shap_interpretability.png"
    if shap_path.exists():
        figs["shap"] = shap_path
    build_docx(figs)
    print(DOCX_OUT)
    for name, path in figs.items():
        print(name, path)


if __name__ == "__main__":
    main()
