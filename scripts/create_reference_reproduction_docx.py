from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from PIL import Image


ROOT = Path(__file__).resolve().parents[1]
FIG_DIR = ROOT / "outputs" / "reference_style_reproduction"
OUT = ROOT / "outputs" / "参考文献图表复现_NSCLC_MMFT.docx"


FIGURES = [
    (
        "Figure 1 | 性能评估复现",
        "对应参考文献 Figure 1：ROC、交叉验证、混淆矩阵、风险分布和阈值指标。ROC/混淆矩阵/风险分布使用本研究公共 PET/CT-blood 测试集真实结果；外部验证框保留待本院队列替换。",
        FIG_DIR / "Figure1_reference_style_performance.png",
    ),
    (
        "Figure 2 | 亚组分析复现",
        "对应参考文献 Figure 2：按年龄、CEA、SUVmax、NLR、分期、性别等简单变量分层，展示 Ridge 与 MMFT 的 AUC 及 bootstrap 95% 区间。为公共测试集探索性亚组分析。",
        FIG_DIR / "Figure2_reference_style_subgroup_auc.png",
    ),
    (
        "Figure 3 | 个体解释性复现",
        "对应参考文献 Figure 3：原文为 Grad-CAM + 风险分布；本研究暂无原始影像 Grad-CAM，因此用 SHAP 个体条形图替代，展示高/低风险代表病例与模型风险分布。",
        FIG_DIR / "Figure3_reference_style_case_interpretability.png",
    ),
    (
        "Figure 4 | 外部 TCGA 生存图",
        "对应参考文献 Figure 4：原文为 KM 曲线。本研究用 UCSC Xena GDC TCGA-LUAD/LUSC 的真实 survival 数据，按 CEACAM5、SLC2A1、HK2、LDHA、TP63、KRT5 构成的 biological risk proxy 中位数分组绘制。",
        FIG_DIR / "Figure4_reference_style_km_tcga_real.png",
    ),
    (
        "Figure 5 | 临床路径复现",
        "对应参考文献 Figure 5：将 HOPE-guided clinical pathway 改写为 NSCLC-MMFT 研究路径，强调术前 PET/CT + 血液/CEA -> 模型输出 -> 本院外部验证。",
        FIG_DIR / "Figure5_reference_style_clinical_pathway.png",
    ),
    (
        "Figure 6 | 纳排/队列流程复现",
        "对应参考文献 Figure 6：展示公共 PET/CT-blood 队列、训练/测试分支、TCGA 干实验分支和本院外部验证分支。",
        FIG_DIR / "Figure6_reference_style_participant_flow.png",
    ),
    (
        "Figure 7 | GPTimage2 AI 生成流程图",
        "对应参考文献 Figure 7：使用 GPTimage2 生成医学 AI workflow 底图，再用程序叠加准确标签，避免 AI 小字错误。推荐作为论文 graphical workflow 初稿。",
        FIG_DIR / "Figure7_GPTimage2_workflow_labeled.png",
    ),
    (
        "Figure 8 | TCGA 外部生物信息学验证",
        "作为新增干实验图：真实 TCGA-LUAD/LUSC Xena 数据，包含 marker 差异散点、关键基因箱线图、通路 proxy 热图和 biological risk proxy KM 曲线。",
        FIG_DIR / "Figure8_TCGA_external_bioinfo_validation.png",
    ),
]

TABLES = [
    (
        "Table 1 | 公共队列基线特征表",
        "对应参考文献 Table 1：使用 PLOS 2024 PET/CT-blood 公共队列真实 baseline 数据，按 histology 0/1 分组。",
        FIG_DIR / "Table1_baseline_characteristics_public_cohort.png",
    ),
    (
        "Table 2 | 模型性能表",
        "对应参考文献 Table 2：使用 Ridge logistic、CV MMFT ensemble 和 MMFT rank-refit 的真实公共测试集结果。",
        FIG_DIR / "Table2_model_performance_public_test.png",
    ),
    (
        "Table 3 | 消融/落地计划表",
        "对应参考文献 Table 3：结合真实基线、已有 MMFT 结果和后续本院验证计划，形成论文方法学落地图。",
        FIG_DIR / "Table3_ablation_and_landing_plan.png",
    ),
]


def set_landscape(section) -> None:
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    section.top_margin = Inches(0.45)
    section.bottom_margin = Inches(0.45)
    section.left_margin = Inches(0.5)
    section.right_margin = Inches(0.5)


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(18 if level == 1 else 13)
    run.font.color.rgb = RGBColor(31, 78, 121)
    p.space_before = Pt(8)
    p.space_after = Pt(5)


def add_note(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(80, 80, 80)
    p.space_after = Pt(4)


def add_image_block(doc: Document, title: str, note: str, path: Path) -> None:
    add_heading(doc, title, level=2)
    add_note(doc, note)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    if path.exists():
        width_inches = 9.4
        try:
            with Image.open(path) as im:
                ratio = im.height / im.width
            width_inches = min(width_inches, 6.15 / ratio)
            width_inches = max(width_inches, 6.8)
        except Exception:
            pass
        run.add_picture(str(path), width=Inches(width_inches))
    else:
        run.add_text(f"Missing image: {path}")
    p.space_after = Pt(10)


def add_mapping_table(doc: Document) -> None:
    rows = [
        ("Fig.1", "ROC + confusion + score distribution", "Figure1", "真实公共测试集 + 待本院外部验证"),
        ("Fig.2", "Subgroup AUC", "Figure2", "真实公共测试集探索性亚组"),
        ("Fig.3", "Grad-CAM + score distribution", "Figure3", "改为 SHAP 个体解释 + 风险分布"),
        ("Fig.4", "Kaplan-Meier", "Figure4", "真实 TCGA survival + biological risk proxy"),
        ("Fig.5", "Clinical pathway", "Figure5", "研究路径图"),
        ("Fig.6", "Participant flow", "Figure6", "公共队列 + TCGA + 本院验证流程"),
        ("Fig.7", "Overall workflow", "Figure7", "GPTimage2 AI 底图 + 程序准确标签"),
        ("Table1", "Baseline characteristics", "Table1", "真实公共 baseline 数据"),
        ("Table2", "Model performance", "Table2", "真实模型结果"),
        ("Table3", "Ablation", "Table3", "真实/已有结果 + 计划项"),
    ]
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    headers = ["参考文献图表", "原文图表功能", "本研究复现", "数据状态"]
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = h
        for p in cell.paragraphs:
            for r in p.runs:
                r.bold = True
                r.font.size = Pt(9)
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = value
            for p in cells[i].paragraphs:
                for r in p.runs:
                    r.font.size = Pt(8.5)


def main() -> None:
    doc = Document()
    set_landscape(doc.sections[0])
    styles = doc.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(9.5)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("参考文献图表复现与本研究对应版")
    r.bold = True
    r.font.size = Pt(20)
    r.font.color.rgb = RGBColor(31, 78, 121)

    add_note(
        doc,
        "目的：尽可能复现 s41746-026-02834-9_reference.pdf 的 Figure 1-7 与 Table 1-3 的图表结构，并映射到 NSCLC-MMFT 研究。"
        "其中模型图使用真实公共 PET/CT-blood 结果，生信/KM 使用真实 TCGA-LUAD/LUSC 外部数据，GPTimage2 用于生成总流程图底图。",
    )
    add_mapping_table(doc)

    for item in FIGURES:
        doc.add_page_break()
        add_image_block(doc, *item)
    for item in TABLES:
        doc.add_page_break()
        add_image_block(doc, *item)

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
