from __future__ import annotations

from pathlib import Path
from shutil import copyfile

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image


ROOT = Path(__file__).resolve().parents[1]
REF = ROOT / "outputs" / "reference_style_reproduction"
GPT = ROOT / "outputs" / "gptimage2_diagrams"
SUPP = ROOT / "outputs" / "supplementary_figures"

OUT_CN = ROOT / "outputs" / "图表展示过程版_NSCLC_MMFT.docx"
OUT_ASCII = ROOT / "outputs" / "student_presentation_figure_plan_nsclc_mmft.docx"

FONT = "Microsoft YaHei"
INK = RGBColor(31, 41, 55)
BLUE = RGBColor(31, 78, 121)
MUTED = RGBColor(85, 96, 110)
RED = RGBColor(153, 27, 27)


FIGURE_ITEMS = [
    {
        "label": "图1",
        "title": "模型总体性能",
        "path": REF / "Figure1_reference_style_performance.png",
        "text": (
            "这张图主要是想先把目前公共队列建模的整体结果展示出来。左上角是 ROC 曲线，"
            "可以看到 MMFT rank-refit 模型在公共测试集上的 AUC 是 0.868，略高于 Ridge logistic 的 0.854，"
            "说明融合模型相比普通线性模型有一定提升。中间的混淆矩阵主要是看模型具体分错了哪些病例，"
            "能补充 AUC 之外的信息。下方风险分布图可以看到，两类患者的预测概率有一定区分，说明模型不是只在统计指标上好看，"
            "而是能够把一部分患者分到相对高风险和低风险区间。右侧目前预留了本院外部验证的位置，后面等本院病例收集完成后，"
            "可以把外部验证 ROC 和混淆矩阵补进去。整体来说，这张图是文章最核心的模型性能图，但目前仍然代表公共数据内部测试结果，"
            "最终可信度需要本院验证来补强。"
        ),
    },
    {
        "label": "图2",
        "title": "亚组稳定性",
        "path": REF / "Figure2_reference_style_subgroup_auc.png",
        "text": (
            "这张图主要是想看模型在不同患者分层里是否还保持稳定。我把年龄、CEA、SUVmax、NLR、分期和性别这些比较容易解释的临床变量分开来看，"
            "每个小图对应一个分层变量，柱子的高度代表 AUC，误差线代表不确定性。多数亚组里 MMFT 仍能维持相对可接受的区分度，"
            "说明模型不是完全依赖某一个单独变量。需要注意的是，有些亚组样本量比较少，比如高分期亚组，置信区间会比较宽，"
            "所以这张图更适合作为探索性稳定性分析。后面本院验证时，我会优先保留样本量够、临床意义明确的几个亚组，不会把亚组切得太碎。"
        ),
    },
    {
        "label": "图3",
        "title": "个体解释与 SHAP",
        "path": REF / "Figure3_reference_style_case_interpretability.png",
        "text": (
            "这张图主要是想解释模型为什么会给出某个风险判断。左右两侧是代表性病例的 SHAP 分解，可以看到哪些变量把预测概率往高风险方向推，"
            "哪些变量把结果往低风险方向拉。中间是整体风险分布，用来观察不同类别患者的预测概率是否有分离趋势。"
            "目前 SHAP 重要性排在前面的主要是融合风险 token、SUVmin、CEA、SUVmean 等，这和我们想讲的影像代谢负荷、血清标志物以及多模态融合方向是吻合的。"
            "这部分的意义不是证明这些变量有因果作用，而是让模型结果更容易解释。后续本院验证完成后，可以再做一次 SHAP，观察变量排序是否一致。"
        ),
    },
    {
        "label": "图4",
        "title": "TCGA 生存曲线",
        "path": REF / "Figure4_reference_style_km_tcga_real.png",
        "text": (
            "这张图是干实验部分，用来给模型变量提供外部生物学解释。这里不是直接验证我们的预测模型，而是用 TCGA-LUAD/LUSC 的表达和生存数据，"
            "把 CEACAM5、SLC2A1、HK2、LDHA、TP63、KRT5 这些和 CEA、糖酵解、鳞癌分化相关的基因合成一个 biological risk proxy。"
            "从图上看，LUAD 队列里高低风险组的生存曲线分离更明显，说明这条生物学线索在外部转录组数据中有一定支持。"
            "这张图后面在文章里更适合写成机制层面的补充证据，而不是写成模型已经能够预测生存。"
        ),
    },
    {
        "label": "图5",
        "title": "临床路径图",
        "path": REF / "Figure5_reference_style_clinical_pathway.png",
        "text": (
            "这张图主要是把模型放回胸外科术前评估的实际场景里。左边是术前比较容易获得的资料，包括影像、血液炎症指标和 CEA，"
            "中间是模型输出的风险概率，右边是后续本院验证和潜在临床分层。这样设计的目的，是说明这个课题不是单纯为了做一个算法，"
            "而是希望围绕术前信息做一个可以被验证的辅助评估流程。现阶段这张图的定位还是研究路径图，后续是否有真实临床应用价值，"
            "要看本院外部验证结果。"
        ),
    },
    {
        "label": "图6",
        "title": "队列与纳排流程",
        "path": REF / "Figure6_reference_style_participant_flow.png",
        "text": (
            "这张图主要是交代整个研究的数据来源和分工。公共 PET/CT-blood 队列负责前期建模和内部测试，TCGA-LUAD/LUSC 队列负责生物信息学论证，"
            "本院约 100 例病例计划作为锁定模型后的独立验证集。这样的设计对毕业课题比较稳：公共数据先保证模型能跑出结果，"
            "TCGA 补充解释性，本院数据再补外部验证。后面本院数据最好只用于验证，不再参与模型调参，这样外部验证结果才更干净。"
        ),
    },
    {
        "label": "图7",
        "title": "研究总流程图",
        "path": GPT / "Figure7_GPTimage2_research_workflow_final.png",
        "text": (
            "这张图是整篇文章的总流程图，用来把数据、模型和分析三部分串起来。上面一行是数据来源，包括胸部 CT/PET、血液和 CEA、本院验证队列以及 TCGA；"
            "中间一行是建模流程，从特征表、缺失值处理和标准化，到 MMFT rank-refit 模型，再到风险概率和 SHAP 解释；"
            "下面一行是结果评价，包括 ROC/AUC、亚组分析、外部生物学论证和临床路径。图的底图是 GPTimage2 生成的，但所有关键文字和节点都是后期程序叠加校正的，"
            "所以可以避免 AI 生图常见的小字错误。正式文章里，这张图可以放在方法部分开头，帮助读者快速理解整体技术路线。"
        ),
    },
    {
        "label": "图8",
        "title": "模型架构图",
        "path": GPT / "Figure9_GPTimage2_model_architecture_final.png",
        "text": (
            "这张图主要是把 MMFT 模型结构说明白。左侧是不同输入分支，包括临床信息、血液和 CEA 指标、PET/CT 代谢特征，以及可选的 CT 语义或 radiomics 分支。"
            "每一类变量先做缺失值处理和标准化，再转换成 token 和 embedding，之后进入中间的 Transformer encoder 学习不同模态之间的交互。"
            "底部的 risk token 用来汇总患者层面的融合风险表示，右侧输出包括 rank-refit 校准、患者风险概率、ROC/AUC、校准/DCA 和 SHAP 解释。"
            "如果后面本院数据没有 PET/CT 指标，这张架构图也可以调整，把 PET/CT 分支作为可选模块，另做胸部 CT 加血液指标的简化模型。"
        ),
    },
    {
        "label": "图9",
        "title": "TCGA 生物信息学补强",
        "path": REF / "Figure8_TCGA_external_bioinfo_validation.png",
        "text": (
            "这张图是对模型变量的生物信息学补强。它主要把 CEA/CEACAM5、糖酵解相关基因以及鳞癌分化标志物放到 TCGA-LUAD/LUSC 里看方向。"
            "A、B 面板展示关键基因在 LUAD 和 LUSC 之间的表达差异，C 面板把这些基因按通路或生物学功能做成 proxy heatmap，"
            "D 面板再用 biological risk proxy 观察生存分层。这样可以把模型中的 CEA 和代谢特征，与外部转录组中的肿瘤分泌表型、糖酵解和病理分化联系起来。"
            "这部分适合写成文章的新颖性补充，但主线仍然应该放在术前多模态模型构建和本院验证上。"
        ),
    },
]


TABLE_ITEMS = [
    {
        "label": "表1",
        "title": "公共队列基线特征",
        "path": REF / "Table1_baseline_characteristics_public_cohort.png",
        "text": (
            "这张表主要是说明公共队列里有哪些基础变量可以用。可以看到，年龄、BMI、血常规、炎症指标、CEA 和 PET 代谢参数都比较完整，"
            "和我后续想在本院收集的术前指标基本能对应上。后面做本院数据表时，我会尽量沿用这里的变量名称和单位，"
            "这样公共模型和本院验证之间更容易衔接。表里的 Class 0/1 在正式方法部分需要写清楚对应的病理标签，避免读者对结局编码产生疑问。"
        ),
    },
    {
        "label": "表2",
        "title": "模型性能表",
        "path": REF / "Table2_model_performance_public_test.png",
        "text": (
            "这张表是把主要模型性能用数字列出来，方便后面写结果部分。Ridge logistic 是传统机器学习强基线，AUC 为 0.854；"
            "MMFT rank-refit 是目前效果最好的模型，AUC 为 0.868，同时 accuracy、specificity 和 F1 也有一定优势。"
            "这个提升不是非常夸张，所以后面文章里不适合只强调深度学习多强，而应该把它写成在强基线基础上的小幅增益，并结合 SHAP、校准和 DCA 来体现完整性。"
        ),
    },
    {
        "label": "表3",
        "title": "消融与落地计划",
        "path": REF / "Table3_ablation_and_landing_plan.png",
        "text": (
            "这张表主要是说明为什么要做融合模型，以及每个模块大概贡献在哪里。表里把 clinical/blood、PET/CT metabolic、CEA、risk token、rank-refit 分开列出来，"
            "可以看出最终模型比传统基线有一定提升，但提升幅度有限。这个结果反而比较真实，也提醒后面写作时要把重点放在可解释性、外部验证和低成本可复现流程上。"
            "本院外部验证现在还是 pending，等数据补齐后再把这一行替换成真实结果。"
        ),
    },
]


SUPPLEMENTARY_ITEMS = [
    {
        "label": "FigS1",
        "title": "数据预处理与缺失值",
        "path": SUPP / "FigS1_data_preprocessing_missingness.png",
        "text": (
            "这张补充图主要是把数据处理过程交代清楚。左上角是关键变量缺失率，公共队列里的临床、血液和 PET 代谢变量缺失率较低；"
            "右上角是特征筛选流程，强调缺失值处理、标准化和特征筛选都应该在训练流程内部完成，避免测试集信息泄漏。"
            "下方展示训练/测试划分和不同特征块的变量数量。后面本院数据接入后，我需要重新画本院队列的缺失率图，因为本院真实病历数据大概率会比公共队列更不整齐。"
        ),
    },
    {
        "label": "FigS2",
        "title": "模型消融实验",
        "path": SUPP / "FigS2_model_ablation.png",
        "text": (
            "这张补充图主要是比较不同特征模块和建模策略。左侧看不同特征块的传统模型表现，右侧把 clinical-blood、PET metabolic、radiomics、risk token、rank-refit 等配置放在一起比较。"
            "它的作用是说明最终模型不是单纯堆复杂结构，而是每一步都有对照。需要注意的是，这里面既有真实 benchmark，也有 prior run 和 final run，"
            "正式成稿时要在图注或方法里把来源说清楚。"
        ),
    },
    {
        "label": "FigS3",
        "title": "TCGA/GEO 关键基因方向热图",
        "path": SUPP / "FigS3_TCGA_GEO_key_gene_direction_heatmap.png",
        "text": (
            "这张补充图是干实验的扩展版，主要看关键基因方向是否和模型解释一致。当前版本基于真实 TCGA-LUAD/LUSC 数据，"
            "可以看到腺癌标志物、鳞癌标志物和糖酵解相关基因在两个病理亚型之间存在方向性差异。GEO 队列的位置先作为扩展槽保留，"
            "等后面确定合适队列并完成下载注释后再补。GEO 没补上之前，正式图题应以 TCGA external validation 为主。"
        ),
    },
    {
        "label": "FigS4",
        "title": "校准曲线与 DCA",
        "path": SUPP / "FigS4_calibration_DCA.png",
        "text": (
            "这张补充图主要是补足模型评价中的校准度和临床净获益。左侧校准曲线看预测概率和真实比例是否接近，右侧 DCA 看不同阈值下模型相比 treat all 和 treat none 是否有净获益。"
            "这部分可以让文章不只是报告 AUC，而是更接近完整的临床预测模型评价。公共测试集样本量还不大，所以目前只能作为初步评价；"
            "本院验证完成后，这张图需要重新绘制并作为重点补充结果。"
        ),
    },
    {
        "label": "FigS5",
        "title": "SHAP 补充图",
        "path": SUPP / "FigS5_SHAP_supplement_composite.png",
        "text": (
            "这张补充图把 SHAP 的几种展示方式放在一起。beeswarm 用来看每个变量对模型输出的方向和离散程度，bar 图看平均绝对贡献，"
            "dependence 图看 CEA 与 SUVmax 相关的局部效应，waterfall 图看代表性病例的预测分解。"
            "这样主文中可以放一张简洁的解释性图，补充材料里再放完整 SHAP 图，既能说明模型不是黑箱，也不会让正文太拥挤。"
        ),
    },
    {
        "label": "FigS6",
        "title": "本院验证队列模板",
        "path": SUPP / "FigS6_hospital_validation_template.png",
        "text": (
            "这张图是本院外部验证的预留版式。后面本院约 100 例病例收齐后，可以在这里放外部验证 ROC、校准曲线、DCA 和亚组稳定性。"
            "它现在的作用是把下一步工作提前框出来，也提醒本院数据只用于锁定模型后的验证。等真实数据接入后，pending 字样需要全部替换成正式结果。"
        ),
    },
]


def set_font(run, size: float | None = None, bold: bool = False, color: RGBColor | None = None) -> None:
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    if size is not None:
        run.font.size = Pt(size)
    run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def setup_section(section) -> None:
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    section.top_margin = Inches(0.42)
    section.bottom_margin = Inches(0.42)
    section.left_margin = Inches(0.48)
    section.right_margin = Inches(0.48)


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    p = doc.add_paragraph()
    r = p.add_run(text)
    set_font(r, 16 if level == 1 else 12.5, True, BLUE if level == 1 else INK)
    p.paragraph_format.space_before = Pt(7 if level == 1 else 4)
    p.paragraph_format.space_after = Pt(4)


def add_para(doc: Document, text: str, size: float = 9.4, color: RGBColor = INK, bold_first: str | None = None) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.16
    p.paragraph_format.space_after = Pt(3)
    if bold_first and text.startswith(bold_first):
        r1 = p.add_run(bold_first)
        set_font(r1, size, True, BLUE)
        r2 = p.add_run(text[len(bold_first):])
        set_font(r2, size, False, color)
    else:
        r = p.add_run(text)
        set_font(r, size, False, color)


def image_width(path: Path, max_width: float = 9.45, max_height: float = 4.65) -> float:
    try:
        with Image.open(path) as im:
            ratio = im.height / im.width
        return max(5.6, min(max_width, max_height / ratio))
    except Exception:
        return max_width


def add_image(doc: Document, path: Path, max_width: float = 9.45, max_height: float = 4.65) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if path.exists():
        p.add_run().add_picture(str(path), width=Inches(image_width(path, max_width, max_height)))
    else:
        r = p.add_run(f"图片缺失：{path}")
        set_font(r, 10, True, RED)
    p.paragraph_format.space_after = Pt(5)


def add_figure_page(doc: Document, item: dict, max_height: float = 4.55) -> None:
    doc.add_page_break()
    add_heading(doc, f"{item['label']}｜{item['title']}", level=1)
    add_image(doc, item["path"], max_height=max_height)
    add_para(doc, "图表解读：" + item["text"], size=9.4, bold_first="图表解读：")


def add_cover(doc: Document) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("NSCLC-MMFT 图表展示过程与结果解读")
    set_font(r, 22, True, BLUE)
    p.paragraph_format.space_after = Pt(3)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("学生汇报视角｜公共队列建模、TCGA 干实验与本院验证预案")
    set_font(r, 11, False, MUTED)
    p.paragraph_format.space_after = Pt(9)

    add_para(
        doc,
        "这版文档主要用于展示我目前对整篇文章图表体系的组织思路。文字部分按照汇报时的说明方式来写，重点是把每张图想表达的意思、当前能支持的结论以及后续需要补充的数据讲清楚。",
        size=10,
    )
    add_para(
        doc,
        "整体研究主线保持为：公共 PET/CT-blood 队列先完成模型构建和内部测试，TCGA-LUAD/LUSC 用于补充生物学解释，本院约 100 例病例后续作为独立验证队列。",
        size=10,
    )

    add_heading(doc, "整体展示顺序", level=1)
    rows = [
        ("第一步", "先用图1-图3展示模型性能、亚组稳定性和可解释性，说明模型目前确实有一定预测信号。"),
        ("第二步", "再用图4和图9展示 TCGA 干实验，说明 CEA、代谢和病理分化相关轴有外部生物学依据。"),
        ("第三步", "用图5-图8说明研究路径、队列分工、总流程和模型结构，保证文章逻辑完整。"),
        ("第四步", "用表1-表3和 FigS1-FigS6 补充数据质量、消融实验、校准、DCA、SHAP 和本院验证预案。"),
    ]
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for i, header in enumerate(["环节", "展示内容"]):
        cell = table.rows[0].cells[i]
        cell.text = header
        shade(cell, "E8EEF5")
        for p in cell.paragraphs:
            for r in p.runs:
                set_font(r, 9, True, BLUE)
    for left, right in rows:
        cells = table.add_row().cells
        cells[0].text = left
        cells[1].text = right
        for i, cell in enumerate(cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cell.paragraphs:
                for r in p.runs:
                    set_font(r, 8.7, i == 0, INK)


def add_index(doc: Document) -> None:
    doc.add_page_break()
    add_heading(doc, "图表清单", level=1)
    add_para(doc, "主图主要承担文章故事线，表格图用于放关键数字，补充图用于补足方法学细节和后续验证计划。", size=9.5)
    rows = []
    for item in FIGURE_ITEMS:
        rows.append((item["label"], item["title"], "主图"))
    for item in TABLE_ITEMS:
        rows.append((item["label"], item["title"], "表格"))
    for item in SUPPLEMENTARY_ITEMS:
        rows.append((item["label"], item["title"], "补充图"))
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for i, header in enumerate(["编号", "标题", "用途"]):
        cell = table.rows[0].cells[i]
        cell.text = header
        shade(cell, "E8EEF5")
        for p in cell.paragraphs:
            for r in p.runs:
                set_font(r, 8.6, True, BLUE)
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = value
            for p in cells[i].paragraphs:
                for r in p.runs:
                    set_font(r, 8.2, False, INK)


def build() -> None:
    doc = Document()
    setup_section(doc.sections[0])
    doc.styles["Normal"].font.name = FONT
    doc.styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    doc.styles["Normal"].font.size = Pt(9.5)

    add_cover(doc)
    add_index(doc)

    for item in FIGURE_ITEMS:
        add_figure_page(doc, item)
    for item in TABLE_ITEMS:
        add_figure_page(doc, item, max_height=4.15)
    for item in SUPPLEMENTARY_ITEMS:
        add_figure_page(doc, item)

    doc.add_page_break()
    add_heading(doc, "下一步需要补齐的内容", level=1)
    add_para(
        doc,
        "目前图表体系已经能支撑阶段性汇报。下一步最关键的是确认最终结局变量。如果继续写病理亚型预测，现有公共库和 TCGA 逻辑比较顺；如果改成淋巴结转移或早期复发，本院数据必须提供对应的病理或随访金标准。",
        size=10,
    )
    add_para(
        doc,
        "本院数据收集时建议提前锁定字段：年龄、性别、吸烟史、血常规、NLR/PLR/SII、CEA、CT 语义特征、PET SUV 指标、最终病理标签和必要的随访信息。这样后面替换 FigS6 和外部验证表会更快。",
        size=10,
    )

    doc.save(OUT_CN)
    copyfile(OUT_CN, OUT_ASCII)
    print(OUT_CN)
    print(OUT_ASCII)


if __name__ == "__main__":
    build()
