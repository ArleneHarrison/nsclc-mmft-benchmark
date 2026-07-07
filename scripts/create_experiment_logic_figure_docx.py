from __future__ import annotations

from pathlib import Path
from shutil import copyfile

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image


ROOT = Path(__file__).resolve().parents[1]
REF = ROOT / "outputs" / "reference_style_reproduction"
GPT = ROOT / "outputs" / "gptimage2_diagrams"
SUPP = ROOT / "outputs" / "supplementary_figures"

OUT_CN = ROOT / "outputs" / "实验思路与图表逻辑总览_NSCLC_MMFT.docx"
OUT_ASCII = ROOT / "outputs" / "experiment_logic_figure_overview_nsclc_mmft.docx"

FONT = "Microsoft YaHei"
INK = RGBColor(31, 41, 55)
BLUE = RGBColor(31, 78, 121)
MUTED = RGBColor(80, 88, 100)
GREEN = RGBColor(40, 93, 56)
RED = RGBColor(153, 27, 27)


CORE_FIGURES = [
    {
        "label": "图7",
        "title": "研究总流程图",
        "path": GPT / "Figure7_GPTimage2_research_workflow_final.png",
        "where": "建议放在方法部分开头，作为整篇文章的图形摘要或技术路线图。",
        "logic": (
            "这张图先把研究讲成一条线：上面是数据来源，中间是模型构建，下面是结果评价。"
            "我这篇文章不是单纯做一个深度学习模型，而是先用公共队列完成可运行的模型，"
            "再用 TCGA 做生物学解释，最后预留本院独立验证。"
        ),
    },
    {
        "label": "图6",
        "title": "队列与纳排流程",
        "path": REF / "Figure6_reference_style_participant_flow.png",
        "where": "建议放在方法部分，紧跟研究总流程图之后。",
        "logic": (
            "这张图回答数据从哪里来、每个队列承担什么任务。公共 PET/CT-blood 队列用于建模和内部测试；"
            "TCGA-LUAD/LUSC 用于干实验论证；本院约 100 例病例后续只做锁定模型后的外部验证。"
        ),
    },
    {
        "label": "FigS1",
        "title": "数据预处理、缺失值与划分",
        "path": SUPP / "FigS1_data_preprocessing_missingness.png",
        "where": "建议放在补充材料，用来说明数据质量和建模前处理。",
        "logic": (
            "这张图把数据清洗过程说清楚：先看缺失率，再做缺失值处理、标准化、特征筛选，"
            "最后进行分层训练/测试划分。这样可以解释模型不是直接把表格丢进去跑，而是有基本的数据管理流程。"
        ),
    },
    {
        "label": "表1",
        "title": "公共队列基线特征",
        "path": REF / "Table1_baseline_characteristics_public_cohort.png",
        "where": "建议放在结果部分第一张表。",
        "logic": (
            "这张表证明公共库里确实有适合本课题的术前变量，包括年龄、血常规、炎症指标、CEA 和 PET 代谢参数。"
            "后面本院数据收集时，也应该尽量按照这个表的变量格式去建表。"
        ),
    },
    {
        "label": "图8",
        "title": "MMFT 模型架构图",
        "path": GPT / "Figure9_GPTimage2_model_architecture_final.png",
        "where": "建议放在方法部分的模型构建小节。",
        "logic": (
            "这张图解释我设计的模型怎么工作。左边是临床、血液/CEA、PET/CT 代谢和可选 CT 语义分支；"
            "中间是 Transformer encoder 学习不同模态之间的关系；右边输出风险概率、校准、ROC、DCA 和 SHAP。"
        ),
    },
    {
        "label": "图1",
        "title": "模型总体性能",
        "path": REF / "Figure1_reference_style_performance.png",
        "where": "建议作为结果部分第一张主图。",
        "logic": (
            "这张图回答模型到底有没有预测能力。公共测试集中 MMFT rank-refit 的 AUC 为 0.868，"
            "高于 Ridge logistic 的 0.854；混淆矩阵和风险分布进一步展示模型能把一部分患者分到较高或较低风险区间。"
        ),
    },
    {
        "label": "表2",
        "title": "模型性能表",
        "path": REF / "Table2_model_performance_public_test.png",
        "where": "建议放在图1之后，用数字补充图1。",
        "logic": (
            "这张表把 AUC、accuracy、sensitivity、specificity、PPV、NPV 和 F1 都列出来，"
            "避免文章只报一个 AUC。它能说明最终模型相比传统基线有小幅提升。"
        ),
    },
    {
        "label": "表3",
        "title": "消融与落地计划",
        "path": REF / "Table3_ablation_and_landing_plan.png",
        "where": "建议放在结果或补充材料，用来解释每个模块的作用。",
        "logic": (
            "这张表说明最终模型不是盲目堆复杂结构，而是把 clinical/blood、PET metabolic、CEA、risk token 和 rank-refit 分开比较。"
            "本院外部验证目前仍是 pending，后面有真实数据后再替换。"
        ),
    },
    {
        "label": "FigS2",
        "title": "模型消融实验",
        "path": SUPP / "FigS2_model_ablation.png",
        "where": "建议放在补充材料，支撑表3。",
        "logic": (
            "这张图更细地展示不同特征块和建模策略的对比，能回答为什么不只做普通 logistic，"
            "也能说明 radiomics-heavy 模型在当前小样本下不一定更好。"
        ),
    },
    {
        "label": "图2",
        "title": "亚组稳定性",
        "path": REF / "Figure2_reference_style_subgroup_auc.png",
        "where": "建议放在主结果或补充结果。",
        "logic": (
            "这张图看模型在不同亚组是否稳定。按年龄、CEA、SUVmax、NLR、分期和性别分层后，"
            "多数亚组仍有一定 AUC，说明模型不是只在某一类患者里有效。样本量较小的亚组只作为探索性结果。"
        ),
    },
    {
        "label": "FigS4",
        "title": "校准曲线与 DCA",
        "path": SUPP / "FigS4_calibration_DCA.png",
        "where": "建议放在补充材料；如果本院验证完成，可以提升为主图。",
        "logic": (
            "这张图补足临床预测模型常见的评价维度。校准曲线看预测概率是否接近真实比例，"
            "DCA 看不同阈值下模型是否有净获益。后续本院验证完成后，这张图需要重新画。"
        ),
    },
    {
        "label": "图3",
        "title": "个体解释与 SHAP",
        "path": REF / "Figure3_reference_style_case_interpretability.png",
        "where": "建议放在模型解释性结果部分。",
        "logic": (
            "这张图解释模型为什么给出某个预测。当前 SHAP 排名前面的变量包括融合风险 token、SUVmin、CEA、SUVmean 等，"
            "和代谢负荷、血清标志物、多模态融合的研究假设一致。"
        ),
    },
    {
        "label": "FigS5",
        "title": "SHAP 补充图",
        "path": SUPP / "FigS5_SHAP_supplement_composite.png",
        "where": "建议放在补充材料，用来完整展示 SHAP 结果。",
        "logic": (
            "这张图把 beeswarm、bar、dependence 和 waterfall 放在一起。它的作用是说明模型不是完全黑箱，"
            "既能看总体变量重要性，也能看单个病例的预测分解。"
        ),
    },
    {
        "label": "图9",
        "title": "TCGA 外部生物信息学验证",
        "path": REF / "Figure8_TCGA_external_bioinfo_validation.png",
        "where": "建议放在干实验结果部分。",
        "logic": (
            "这张图把模型变量背后的生物学线索放到 TCGA 中验证。CEA/CEACAM5、糖酵解相关基因、腺癌和鳞癌标志物"
            "在 LUAD 与 LUSC 之间呈现方向性差异，能支撑模型变量具有一定生物学合理性。"
        ),
    },
    {
        "label": "图4",
        "title": "TCGA 生存曲线",
        "path": REF / "Figure4_reference_style_km_tcga_real.png",
        "where": "建议放在干实验结果部分，紧跟图9。",
        "logic": (
            "这张图不是验证模型本身，而是用 TCGA survival 数据看 biological risk proxy 是否和预后有关。"
            "LUAD 中曲线分离更明显，说明 CEA/代谢/病理分化这条线具有一定外部生物学支持。"
        ),
    },
    {
        "label": "FigS3",
        "title": "关键基因方向热图",
        "path": SUPP / "FigS3_TCGA_GEO_key_gene_direction_heatmap.png",
        "where": "建议放在补充材料，作为 TCGA 干实验的扩展图。",
        "logic": (
            "这张图把关键基因方向更集中地展示出来。目前 TCGA 部分是真实数据，GEO 位置作为后续扩展预留。"
            "在 GEO 没正式接入前，正式图题建议写 TCGA external validation。"
        ),
    },
    {
        "label": "图5",
        "title": "临床路径图",
        "path": REF / "Figure5_reference_style_clinical_pathway.png",
        "where": "建议放在讨论或临床转化设想部分。",
        "logic": (
            "这张图把模型放回胸外科术前评估场景里。它说明模型不是为了替代医生，而是作为术前风险提示和后续验证工具。"
            "真正能否落地，还要看本院验证结果。"
        ),
    },
    {
        "label": "FigS6",
        "title": "本院验证队列模板",
        "path": SUPP / "FigS6_hospital_validation_template.png",
        "where": "目前作为补充图模板；本院数据完成后替换为正式外部验证图。",
        "logic": (
            "这张图是下一阶段工作的占位图。后面本院约 100 例病例完成后，这里会补入外部验证 ROC、校准、DCA 和亚组稳定性。"
            "目前必须保留 pending 标注，不能写成已完成结果。"
        ),
    },
]


BACKUP_FIGURES = [
    {
        "label": "备用图A",
        "title": "补充图总览样式图",
        "path": SUPP / "Supplementary_Figure_Style_Sheet_NSCLC_MMFT.png",
        "logic": "这是给阶段性汇报看的总览拼图，和 FigS1-FigS6 单图内容重复，正式文档中不建议作为独立论文图。",
    },
    {
        "label": "备用图B",
        "title": "程序绘制版流程图",
        "path": REF / "Figure7_reference_style_workflow_programmatic.png",
        "logic": "这是不用 AI 底图的程序版流程图。如果对 GPTimage2 风格不满意，可以用这张替代图7。",
    },
    {
        "label": "备用图C",
        "title": "旧版 GPTimage2 流程图",
        "path": REF / "Figure7_GPTimage2_workflow_labeled.png",
        "logic": "这是早期 GPTimage2 版本，已经被新版图7替代。保留在附录中只是为了说明生成过程，没有必要放入正式稿。",
    },
    {
        "label": "备用图D",
        "title": "KM 模板图",
        "path": REF / "Figure4_reference_style_km_template.png",
        "logic": "这是早期没有 TCGA survival 数据时做的模板图，已经被真实 TCGA KM 图替代，不建议用于正式论文。",
    },
]


SHAP_SINGLE = [
    ("SHAP单图A", "beeswarm", SUPP / "FigS5A_SHAP_beeswarm_highres.png"),
    ("SHAP单图B", "bar", SUPP / "FigS5B_SHAP_bar_highres.png"),
    ("SHAP单图C", "dependence", SUPP / "FigS5C_SHAP_dependence_CEA_highres.png"),
    ("SHAP单图D", "waterfall", SUPP / "FigS5D_SHAP_waterfall_highres.png"),
]


def set_font(run, size: float | None = None, bold: bool = False, color: RGBColor | None = None) -> None:
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    if size is not None:
        run.font.size = Pt(size)
    run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def setup_section(section) -> None:
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    section.top_margin = Inches(0.42)
    section.bottom_margin = Inches(0.42)
    section.left_margin = Inches(0.48)
    section.right_margin = Inches(0.48)


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    p = doc.add_paragraph()
    r = p.add_run(text)
    set_font(r, 16 if level == 1 else 12.5, True, BLUE if level == 1 else INK)
    p.paragraph_format.space_before = Pt(7 if level == 1 else 4)
    p.paragraph_format.space_after = Pt(4)


def add_para(doc: Document, text: str, size: float = 9.3, color: RGBColor = INK) -> None:
    p = doc.add_paragraph()
    r = p.add_run(text)
    set_font(r, size, False, color)
    p.paragraph_format.line_spacing = 1.16
    p.paragraph_format.space_after = Pt(3)


def add_labeled(doc: Document, label: str, text: str, color: RGBColor = BLUE) -> None:
    p = doc.add_paragraph()
    r1 = p.add_run(label)
    set_font(r1, 9.1, True, color)
    r2 = p.add_run(text)
    set_font(r2, 9.1, False, INK)
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(3)


def image_width(path: Path, max_width: float = 9.45, max_height: float = 4.55) -> float:
    try:
        with Image.open(path) as im:
            ratio = im.height / im.width
        return max(5.5, min(max_width, max_height / ratio))
    except Exception:
        return max_width


def add_image(doc: Document, path: Path, max_width: float = 9.45, max_height: float = 4.55) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if path.exists():
        p.add_run().add_picture(str(path), width=Inches(image_width(path, max_width, max_height)))
    else:
        r = p.add_run(f"图片缺失：{path}")
        set_font(r, 10, True, RED)
    p.paragraph_format.space_after = Pt(5)


def add_cover(doc: Document) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("NSCLC-MMFT 实验思路与图表逻辑总览")
    set_font(r, 22, True, BLUE)
    p.paragraph_format.space_after = Pt(3)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("从数据发现、模型构建、基因筛选到图表安排")
    set_font(r, 11, False, MUTED)
    p.paragraph_format.space_after = Pt(10)

    add_para(
        doc,
        "一句话概括：我先找到一个能提供术前 PET/CT、血液指标和 CEA 的公共 NSCLC 队列，用它完成模型构建；再用 TCGA-LUAD/LUSC 解释模型中 CEA、代谢和病理分化相关信号的生物学来源；最后预留本院约 100 例病例作为独立验证。",
        size=10.2,
    )
    add_para(
        doc,
        "这份文档按实验推进顺序整理，不再按单张图零散解释。每张图都会说明它在整条研究逻辑中回答什么问题、用了什么方法、适合放在文章哪里。",
        size=10.2,
    )


def add_story(doc: Document) -> None:
    doc.add_page_break()
    add_heading(doc, "一、整个实验是怎么想出来的", 1)
    steps = [
        ("1. 先从现实限制出发", "本院数据大概只有 100 例左右，如果一开始就用本院数据训练复杂深度学习模型，样本量不够，结果也不稳。因此先用公共数据库完成前期建模。"),
        ("2. 找到可运行的公共队列", "公共 PET/CT-blood 队列有 255 例 NSCLC，包含临床信息、血常规/炎症指标、CEA、PET 代谢参数和 CT/PET radiomics，并且有病理亚型标签。"),
        ("3. 先建立强基线", "先跑 Ridge logistic、LASSO、SVM、随机森林等传统模型，确认简单模型已经能达到较好的 AUC。这样后面深度学习模型才有比较对象。"),
        ("4. 再设计 MMFT 模型", "把临床/血液、PET/CT 代谢和可选 CT 语义特征作为不同输入分支，用 Transformer encoder 学习模态间关系，再用 risk token 汇总患者层面的风险信息。"),
        ("5. 用 SHAP 解释模型", "模型跑出结果后，用 SHAP 看哪些变量在推动预测。当前重要变量集中在 risk token、SUVmin、CEA、SUVmean 等，和影像代谢及血清标志物方向一致。"),
        ("6. 用 TCGA 做生物学论证", "模型里出现的 CEA 和代谢特征不能只停留在统计层面，所以再用 TCGA-LUAD/LUSC 看 CEACAM5、糖酵解和病理分化标志物是否有外部表达和生存证据。"),
        ("7. 最后回到本院验证", "公共库负责建模，TCGA 负责解释，本院病例负责验证。后面本院数据不建议再参与调参，而是作为锁定模型后的外部验证。"),
    ]
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for i, h in enumerate(["推进步骤", "具体逻辑"]):
        table.rows[0].cells[i].text = h
        shade(table.rows[0].cells[i], "E8EEF5")
        for p in table.rows[0].cells[i].paragraphs:
            for r in p.runs:
                set_font(r, 8.8, True, BLUE)
    for left, right in steps:
        cells = table.add_row().cells
        cells[0].text = left
        cells[1].text = right
        for i, cell in enumerate(cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cell.paragraphs:
                for r in p.runs:
                    set_font(r, 8.3, i == 0, INK)


def add_methods(doc: Document) -> None:
    doc.add_page_break()
    add_heading(doc, "二、用了哪些方法", 1)
    rows = [
        ("公共数据建模", "合并 baseline、CT radiomics、PET radiomics 表格；结局为 histology 0/1；分层 70/30 训练测试划分。", "Table1, FigS1, Figure6"),
        ("传统模型基线", "Ridge/LASSO logistic、SVM、随机森林、ExtraTrees、XGBoost、LightGBM 等；主要评价 AUC、accuracy、sensitivity、specificity。", "Figure1, Table2, FigS2"),
        ("MMFT 模型", "不同模态先标准化并 token 化，再输入 Transformer encoder；risk token 汇总融合风险；rank-refit 提高小样本排序稳定性。", "Figure8, Figure1, Table3"),
        ("模型解释", "SHAP beeswarm、bar、dependence、waterfall；重点看风险 token、SUV、CEA 等变量贡献。", "Figure3, FigS5"),
        ("TCGA 干实验", "UCSC Xena 下载 TCGA-LUAD/LUSC expression 和 survival；关键基因差异、通路 proxy、KM 生存分析。", "Figure9, Figure4, FigS3"),
        ("后续外部验证", "本院约 100 例病例只做锁定模型后的验证；计划补外部 ROC、校准、DCA、亚组稳定性。", "Figure5, FigS6"),
    ]
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for i, h in enumerate(["模块", "方法", "对应图表"]):
        table.rows[0].cells[i].text = h
        shade(table.rows[0].cells[i], "E8EEF5")
        for p in table.rows[0].cells[i].paragraphs:
            for r in p.runs:
                set_font(r, 8.8, True, BLUE)
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = value
            for p in cells[i].paragraphs:
                for r in p.runs:
                    set_font(r, 8.0, False, INK)


def add_gene_logic(doc: Document) -> None:
    doc.add_page_break()
    add_heading(doc, "三、TCGA 里的基因是怎么筛出来的", 1)
    add_para(
        doc,
        "这里不是做全基因海选，也不是从几万个基因里硬找显著性。我的筛选逻辑是从模型变量反推生物学轴：模型里有 CEA 和 PET 代谢指标，结局又和 NSCLC 病理亚型相关，所以基因选择围绕 CEA/腺癌分泌表型、糖酵解代谢、鳞癌分化、免疫炎症和基质反应这几条线展开。",
        size=9.8,
    )
    rows = [
        ("CEA / 腺癌分泌表型", "CEA 是模型变量之一；CEACAM5 是 CEA 对应的核心基因。", "CEACAM5, NKX2-1, NAPSA, SFTPB, KRT7, KRT18, KRT19"),
        ("PET 代谢 / 糖酵解", "SUVmean、SUVmax、SUVmin 代表肿瘤代谢负荷，因此选择糖酵解和缺氧相关基因。", "SLC2A1, HK2, LDHA, PKM, ENO1, ALDOA, GAPDH, HIF1A, CA9"),
        ("鳞癌分化", "病理亚型是当前公共库结局，需要加入鳞癌分化标志物作为对照轴。", "TP63, KRT5, KRT6A, KRT14, DSG3, SOX2"),
        ("免疫 / 炎症", "血常规和炎症指标提示免疫炎症背景，TCGA 中用免疫细胞和 checkpoint 相关基因补充解释。", "CD274, CD8A, CD8B, CD3D, CD3E, CD68, CD163, S100A8, S100A9, CXCL8"),
        ("基质 / EMT", "肿瘤侵袭和间质反应作为补充生物学背景，不作为主轴。", "VIM, COL1A1, ACTA2, MMP9, TWIST1, SNAI1, ZEB1"),
    ]
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for i, h in enumerate(["基因组", "为什么选", "代表基因"]):
        table.rows[0].cells[i].text = h
        shade(table.rows[0].cells[i], "E8EEF5")
        for p in table.rows[0].cells[i].paragraphs:
            for r in p.runs:
                set_font(r, 8.6, True, BLUE)
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = value
            for p in cells[i].paragraphs:
                for r in p.runs:
                    set_font(r, 7.9, False, INK)
    add_para(
        doc,
        "分析方法上，先比较 LUAD 与 LUSC 的表达差异，使用 Mann-Whitney U 检验并做 FDR 校正；再把同一生物学轴的基因做 z-score 后求平均，形成 pathway proxy；最后把 CEACAM5、糖酵解和鳞癌分化相关信号合成 exploratory biological risk proxy，用中位数分组画 KM 曲线。",
        size=9.4,
    )


def add_figure_index(doc: Document) -> None:
    doc.add_page_break()
    add_heading(doc, "四、所有图表放置总表", 1)
    rows = []
    for fig in CORE_FIGURES:
        rows.append((fig["label"], fig["title"], fig["where"]))
    for fig in BACKUP_FIGURES:
        rows.append((fig["label"], fig["title"], "备用/不建议作为正式主图，见附录说明。"))
    for label, title, _ in SHAP_SINGLE:
        rows.append((label, title, "SHAP 单图备用版，可拆分放补充材料。"))
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for i, h in enumerate(["图表", "名称", "建议位置"]):
        table.rows[0].cells[i].text = h
        shade(table.rows[0].cells[i], "E8EEF5")
        for p in table.rows[0].cells[i].paragraphs:
            for r in p.runs:
                set_font(r, 8.6, True, BLUE)
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = value
            for p in cells[i].paragraphs:
                for r in p.runs:
                    set_font(r, 7.4, False, INK)


def add_figure_page(doc: Document, fig: dict, backup: bool = False) -> None:
    doc.add_page_break()
    add_heading(doc, f"{fig['label']}｜{fig['title']}", 1)
    add_image(doc, fig["path"], max_height=4.35 if backup else 4.55)
    if "where" in fig:
        add_labeled(doc, "放置位置：", fig["where"], BLUE)
    add_labeled(doc, "图表逻辑：", fig["logic"], GREEN if not backup else MUTED)


def add_shap_single_page(doc: Document) -> None:
    doc.add_page_break()
    add_heading(doc, "SHAP 单图备用版", 1)
    add_para(doc, "这些图和 FigS5 合并图内容一致，但如果后续需要把 SHAP 拆开排版，可以直接使用。", size=9.1, color=MUTED)
    table = doc.add_table(rows=2, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for idx, (label, title, path) in enumerate(SHAP_SINGLE):
        cell = table.rows[idx // 2].cells[idx % 2]
        cell.text = ""
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(f"{label}｜{title}")
        set_font(r, 8.2, True, BLUE)
        if path.exists():
            p_img = cell.add_paragraph()
            p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_img.add_run().add_picture(str(path), width=Inches(4.6))


def add_final_note(doc: Document) -> None:
    doc.add_page_break()
    add_heading(doc, "七、这篇文章最后要讲成什么样", 1)
    add_para(
        doc,
        "最终文章不要讲成“我做了一个很复杂的 Transformer”。更清楚的讲法是：我基于可获得的术前临床、血液和影像代谢变量，先在公共 NSCLC 队列中构建一个多模态预测模型；然后用 SHAP 解释模型变量；再用 TCGA 从 CEA、糖酵解和病理分化角度补充生物学证据；最后计划用本院病例做独立验证。",
        size=10,
    )
    add_para(
        doc,
        "真正的主线是“公共库建模 + 模型解释 + TCGA 干实验 + 本院验证预案”。本院数据补齐后，最重要的是替换 FigS6，并在图1或补充图中加入真实外部验证 ROC、校准曲线和 DCA。",
        size=10,
    )


def build() -> None:
    doc = Document()
    setup_section(doc.sections[0])
    doc.styles["Normal"].font.name = FONT
    doc.styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    doc.styles["Normal"].font.size = Pt(9.4)

    add_cover(doc)
    add_story(doc)
    add_methods(doc)
    add_gene_logic(doc)
    add_figure_index(doc)

    add_heading(doc, "五、正文/补充图逐图说明", 1)
    for fig in CORE_FIGURES:
        add_figure_page(doc, fig)

    doc.add_page_break()
    add_heading(doc, "六、备用图与重复图说明", 1)
    add_para(doc, "下面这些图没有漏掉，但不一定建议放进正式文章正文。它们主要作为汇报、备用或历史版本保留。", size=9.4, color=MUTED)
    for fig in BACKUP_FIGURES:
        add_figure_page(doc, fig, backup=True)
    add_shap_single_page(doc)
    add_final_note(doc)

    doc.save(OUT_CN)
    copyfile(OUT_CN, OUT_ASCII)
    print(OUT_CN)
    print(OUT_ASCII)


if __name__ == "__main__":
    build()
