"""Reusable Markdown -> .docx renderer with Chinese-font support + figure embedding.

Supports: #..###### headings, paragraphs, **bold**/*italic* inline, GFM pipe tables,
bullet (-, *) and numbered (1.) lists, horizontal rules (---), blockquotes (>),
and GFM images ![caption](path). Every run is set to Microsoft YaHei so Chinese
renders correctly. Import and call render_markdown(doc, text, base_dir).
"""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image

FONT = "Microsoft YaHei"
INK = RGBColor(0x1F, 0x29, 0x37)
BLUE = RGBColor(0x1F, 0x4E, 0x79)
MUTED = RGBColor(0x50, 0x58, 0x64)
HEAD_COLORS = {1: RGBColor(0x0B, 0x35, 0x5E), 2: BLUE, 3: RGBColor(0x2A, 0x52, 0x6E),
               4: MUTED, 5: MUTED, 6: MUTED}
HEAD_SIZES = {1: 18, 2: 14.5, 3: 12.5, 4: 11.5, 5: 11, 6: 10.5}


def _set_run_font(run, size=10.5, bold=False, italic=False, color=INK):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), FONT)


_INLINE = re.compile(r"(\*\*.+?\*\*|__.+?__|\*[^*]+?\*|`[^`]+?`)")


def _add_inline(paragraph, text, size=10.5, base_bold=False, color=INK):
    text = text.replace("\\n", " ")
    parts = _INLINE.split(text)
    for part in parts:
        if not part:
            continue
        if (part.startswith("**") and part.endswith("**")) or (part.startswith("__") and part.endswith("__")):
            run = paragraph.add_run(part[2:-2]); _set_run_font(run, size, True, False, color)
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            run = paragraph.add_run(part[1:-1]); _set_run_font(run, size, base_bold, True, color)
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1]); _set_run_font(run, size - 0.5, base_bold, False, RGBColor(0xB2, 0x18, 0x2B))
        else:
            run = paragraph.add_run(part); _set_run_font(run, size, base_bold, False, color)


def _shade(cell, hexcolor):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hexcolor)
    tcpr.append(shd)


def _set_cell_margins(cell, top=20, bottom=20, left=60, right=60):
    tcpr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for tag, val in (("top", top), ("bottom", bottom), ("left", left), ("right", right)):
        node = OxmlElement(f"w:{tag}")
        node.set(qn("w:w"), str(val)); node.set(qn("w:type"), "dxa")
        mar.append(node)
    tcpr.append(mar)


def _set_row_height(row, twips=200):
    trpr = row._tr.get_or_add_trPr()
    h = OxmlElement("w:trHeight")
    h.set(qn("w:val"), str(twips)); h.set(qn("w:hRule"), "atLeast")
    trpr.append(h)


def _add_table(doc, rows):
    if not rows:
        return
    ncol = max(len(r) for r in rows)
    rows = [r + [""] * (ncol - len(r)) for r in rows]
    table = doc.add_table(rows=len(rows), cols=ncol)
    table.style = "Light Grid Accent 1"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.allow_autofit = True
    for i, row in enumerate(rows):
        _set_row_height(table.rows[i], twips=170)
        for j, txt in enumerate(row):
            cell = table.cell(i, j)
            cell.text = ""
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            _set_cell_margins(cell)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0); p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            _add_inline(p, txt.strip(), size=8.5, base_bold=(i == 0),
                        color=(RGBColor(0xFF, 0xFF, 0xFF) if i == 0 else INK))
            if i == 0:
                _shade(cell, "1F4E79")
            elif i % 2 == 0:
                _shade(cell, "EEF3F8")
    tail = doc.add_paragraph()
    tail.paragraph_format.space_after = Pt(2)
    tail.paragraph_format.space_before = Pt(0)


def _is_table_sep(line):
    return bool(re.match(r"^\s*\|?[\s:|-]+\|?\s*$", line)) and "-" in line


_PIPE_PLACEHOLDER = "\x00PIPE\x00"


def _split_row(line):
    line = line.strip().replace("\\|", _PIPE_PLACEHOLDER)
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|"):
        line = line[:-1]
    return [c.strip().replace(_PIPE_PLACEHOLDER, "|") for c in line.split("|")]


def _add_code_block(doc, code_lines):
    if not any(l.strip() for l in code_lines):
        return
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.left_indent = Inches(0.15)
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), "F3F5F7")
    pPr.append(shd)
    pbdr = OxmlElement("w:pBdr")
    for edge in ("top", "bottom", "left", "right"):
        b = OxmlElement(f"w:{edge}")
        b.set(qn("w:val"), "single"); b.set(qn("w:sz"), "4"); b.set(qn("w:space"), "4"); b.set(qn("w:color"), "C9D6E3")
        pbdr.append(b)
    pPr.append(pbdr)
    text = "\n".join(code_lines).rstrip("\n")
    for idx, seg in enumerate(text.split("\n")):
        if idx > 0:
            p.add_run().add_break()
        run = p.add_run(seg if seg else " ")
        run.font.name = "Consolas"
        run.font.size = Pt(9)
        run.font.color.rgb = INK
        rpr = run._element.get_or_add_rPr()
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is None:
            rfonts = OxmlElement("w:rFonts"); rpr.append(rfonts)
        for attr in ("w:ascii", "w:hAnsi", "w:cs"):
            rfonts.set(qn(attr), "Consolas")
        rfonts.set(qn("w:eastAsia"), FONT)


def add_figure(doc, path, caption=None, width_in=6.4, takeaway=None):
    path = Path(path)
    if not path.exists():
        p = doc.add_paragraph(); _add_inline(p, f"[缺图: {path.name}]", 9, color=RGBColor(0xB2,0x18,0x2B))
        return
    try:
        with Image.open(path) as im:
            w, h = im.size
        ratio = h / w
    except Exception:
        ratio = 0.6
    width = min(width_in, 6.6)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2); p.paragraph_format.space_before = Pt(2)
    run = p.add_run()
    run.add_picture(str(path), width=Inches(width))
    if caption:
        cp = doc.add_paragraph(); cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cp.paragraph_format.space_after = Pt(2); cp.paragraph_format.space_before = Pt(0)
        _add_inline(cp, caption, size=9, base_bold=True, color=BLUE)
    if takeaway:
        tp = doc.add_paragraph()
        tp.paragraph_format.space_after = Pt(2); tp.paragraph_format.space_before = Pt(0)
        tp.paragraph_format.line_spacing = 1.05
        _add_inline(tp, takeaway, size=9.5, color=MUTED)
    tail = doc.add_paragraph(); tail.paragraph_format.space_after = Pt(2); tail.paragraph_format.space_before = Pt(0)


def render_markdown(doc, text, base_dir="."):
    base_dir = Path(base_dir)
    lines = text.replace("\r\n", "\n").split("\n")
    i = 0
    n = len(lines)
    while i < n:
        line = lines[i]
        stripped = line.strip()

        # table
        if "|" in line and i + 1 < n and _is_table_sep(lines[i + 1]):
            header = _split_row(line)
            rows = [header]
            i += 2
            while i < n and "|" in lines[i] and lines[i].strip():
                rows.append(_split_row(lines[i])); i += 1
            _add_table(doc, rows)
            continue

        if not stripped:
            i += 1
            continue

        # fenced code block
        if stripped.startswith("```"):
            i += 1
            code_lines = []
            while i < n and not lines[i].strip().startswith("```"):
                code_lines.append(lines[i]); i += 1
            i += 1  # skip closing fence
            _add_code_block(doc, code_lines)
            continue

        # horizontal rule
        if re.match(r"^\s*(-{3,}|\*{3,}|_{3,})\s*$", line):
            hr = doc.add_paragraph(); hr.paragraph_format.space_before = Pt(2); hr.paragraph_format.space_after = Pt(2)
            pPr = hr._p.get_or_add_pPr(); pbdr = OxmlElement("w:pBdr"); bottom = OxmlElement("w:bottom")
            bottom.set(qn("w:val"), "single"); bottom.set(qn("w:sz"), "6"); bottom.set(qn("w:space"), "1"); bottom.set(qn("w:color"), "9AA0A6")
            pbdr.append(bottom); pPr.append(pbdr)
            i += 1
            continue

        # image ![cap](path)
        m = re.match(r"^!\[(.*?)\]\((.*?)\)\s*$", stripped)
        if m:
            cap, path = m.group(1), m.group(2)
            fp = Path(path)
            if not fp.is_absolute():
                fp = base_dir / path
            add_figure(doc, fp, caption=cap or None)
            i += 1
            continue

        # heading
        hm = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if hm:
            level = len(hm.group(1)); content = hm.group(2).strip().replace("**", "")
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8 if level <= 2 else 5)
            p.paragraph_format.space_after = Pt(2)
            _add_inline(p, content, size=HEAD_SIZES.get(level, 11), base_bold=True,
                        color=HEAD_COLORS.get(level, INK))
            if level <= 2:
                pPr = p._p.get_or_add_pPr(); pbdr = OxmlElement("w:pBdr"); bottom = OxmlElement("w:bottom")
                bottom.set(qn("w:val"), "single"); bottom.set(qn("w:sz"), "6"); bottom.set(qn("w:space"), "2")
                bottom.set(qn("w:color"), "1F4E79" if level == 1 else "C9D6E3")
                pbdr.append(bottom); pPr.append(pbdr)
            i += 1
            continue

        # blockquote
        if stripped.startswith(">"):
            p = doc.add_paragraph(); p.paragraph_format.left_indent = Inches(0.25)
            _add_inline(p, stripped[1:].strip(), size=10, color=MUTED)
            i += 1
            continue

        # bullet list
        bm = re.match(r"^(\s*)([-*+])\s+(.*)$", line)
        if bm:
            indent = len(bm.group(1))
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.left_indent = Inches(0.25 + 0.25 * (indent // 2))
            p.paragraph_format.space_after = Pt(1)
            _add_inline(p, bm.group(3), size=10.5)
            i += 1
            continue

        # numbered list
        nm = re.match(r"^(\s*)(\d+)[.)]\s+(.*)$", line)
        if nm:
            p = doc.add_paragraph(style="List Number")
            p.paragraph_format.space_after = Pt(1)
            _add_inline(p, nm.group(3), size=10.5)
            i += 1
            continue

        # normal paragraph (merge following non-empty, non-special lines)
        buf = [stripped]
        i += 1
        while i < n and lines[i].strip() and not re.match(r"^(#{1,6}\s|[-*+]\s|\d+[.)]\s|>|!\[)", lines[i].strip()) \
                and "|" not in lines[i] and not re.match(r"^\s*(-{3,})\s*$", lines[i]):
            buf.append(lines[i].strip()); i += 1
        p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.08
        _add_inline(p, " ".join(buf), size=10.5)


def new_document(landscape=False, margin_in=0.7):
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = FONT
    style.font.size = Pt(10.5)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    sec = doc.sections[0]
    if landscape:
        from docx.enum.section import WD_ORIENT
        sec.orientation = WD_ORIENT.LANDSCAPE
        sec.page_width, sec.page_height = sec.page_height, sec.page_width
    for attr in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(sec, attr, Inches(margin_in))
    return doc
